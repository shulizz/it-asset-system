"""Atomic, super-administrator-only creation of users from Excel."""

import re
from io import BytesIO
from urllib.parse import quote

from database import get_db
from deps import require_super_admin
from excel_input import MAX_UPLOAD, read_rows
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from models import OperationLog, User
from permissions import ALL_PERMISSIONS, PERMISSION_KEYS
from pydantic import ValidationError
from sqlalchemy import text
from sqlalchemy.orm import Session

from routers.auth import UserIn, new_user_values, pwd_context

router = APIRouter(prefix="/api/user-import", tags=["user-import"])
HEADERS = ["用户名", "姓名", "初始密码", "邮箱", "角色名称", "部门", "数据范围", "权限"]
MAX_USERS = 100


@router.get("/template")
def template(current_user=Depends(require_super_admin)):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "新增用户"
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H101"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2865DB")
    for column, width in zip("ABCDEFGH", [22, 18, 26, 30, 22, 22, 20, 48]):
        ws.column_dimensions[column].width = width
        for row in range(2, 102):
            ws[f"{column}{row}"].number_format = "@"
    scope = DataValidation(type="list", formula1='"仅本部门,全部部门"')
    ws.add_data_validation(scope)
    scope.add("G2:G101")
    notes = wb.create_sheet("填写说明")
    notes.append(["字段", "填写规则"])
    rules = [
        ("用户名", "必填，最多50字；必须是新账号，不能与系统或本表其他行重复"),
        ("姓名", "必填，最多50字"),
        ("初始密码", "必填，至少10个字符，最多72个UTF-8字节；账号和密码列按文本填写"),
        ("邮箱", "可选，有效邮箱地址，可用于接收提醒"),
        ("角色名称", "可选，仅显示名称，不自动授予权限；超级管理员请在页面手动创建"),
        ("部门", "填写系统中已存在的部门名称，不会自动创建部门"),
        ("数据范围", "仅本部门：必须填写部门；全部部门：可不填部门。空白按仅本部门处理"),
        ("权限", "可选，中文权限名称或英文标识，多项用顿号、逗号或换行分隔；留空表示无操作权限"),
        (
            "导入规则",
            "最多100个用户；全部行检查通过才写入，不覆盖或修改已有用户；只有首个工作表导入",
        ),
    ]
    for item in rules:
        notes.append(item)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 105
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    permissions = wb.create_sheet("权限对照")
    permissions.append(["中文权限名称", "英文标识"])
    for key, label in ALL_PERMISSIONS:
        permissions.append([label, key])
    permissions.column_dimensions["A"].width = 46
    permissions.column_dimensions["B"].width = 24
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''" + quote("新增用户_导入模板.xlsx")
        },
    )


@router.post("/import")
def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user=Depends(require_super_admin),
):
    actor_id, actor_version = current_user.id, current_user.token_version
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "请上传.xlsx用户模板")
    content = file.file.read(MAX_UPLOAD + 1)
    if len(content) > MAX_UPLOAD:
        raise HTTPException(413, "导入文件不能超过10MB")
    rows = read_rows(content, first_sheet=True)
    if not rows:
        raise HTTPException(400, "表格为空")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    if any(header and header not in HEADERS for header in headers):
        raise HTTPException(400, "表头包含不支持的列，请使用用户模板")
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise HTTPException(400, "表头存在重复列")
    if not {"用户名", "姓名", "初始密码"}.issubset(headers):
        raise HTTPException(400, "缺少用户名、姓名或初始密码列，请使用用户模板")
    meaningful = [
        (number, row)
        for number, row in enumerate(rows[1:], 2)
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not meaningful:
        raise HTTPException(400, "表格中没有待导入用户")
    if len(meaningful) > MAX_USERS:
        raise HTTPException(400, "每次最多导入100个用户，请拆分表格")
    labels = {label: key for key, label in ALL_PERMISSIONS}
    seen = set()
    prepared = []
    errors = []
    for number, row in meaningful:
        values = dict(zip(headers, row))

        def value(name):
            return str(values.get(name)).strip() if values.get(name) is not None else ""

        try:
            username = value("用户名")
            if username in seen:
                raise ValueError("用户名在本表中重复")
            seen.add(username)
            role = value("角色名称")
            if role == "super_admin":
                raise ValueError("超级管理员账号请在页面手动创建")
            tokens = [
                part.strip() for part in re.split(r"[,，、;；\n]", value("权限")) if part.strip()
            ]
            grants = [labels.get(part, part) for part in tokens]
            if set(grants) - PERMISSION_KEYS:
                raise ValueError("权限名称无效，请参考权限对照工作表")
            scope = {
                "": "department",
                "仅本部门": "department",
                "全部部门": "all",
                "department": "department",
                "all": "all",
            }.get(value("数据范围"))
            if scope is None:
                raise ValueError("数据范围须填写仅本部门或全部部门")
            # Passwords are kept exactly as entered, without trimming or logging.
            password = str(values["初始密码"]) if values.get("初始密码") is not None else ""
            data = UserIn(
                username=username,
                name=value("姓名"),
                password=password,
                email=value("邮箱") or None,
                role=role,
                department=value("部门") or None,
                data_scope=scope,
                permissions=sorted(set(grants)),
            )
            prepared.append((number, data, new_user_values(data, db)))
        except ValidationError as error:
            errors.append(
                {
                    "row": number,
                    "message": "；".join(item["msg"] for item in error.errors(include_input=False)),
                }
            )
        except HTTPException as error:
            errors.append({"row": number, "message": str(error.detail)})
        except ValueError as error:
            errors.append({"row": number, "message": str(error)})
    if errors:
        raise HTTPException(
            400, {"message": "未导入任何用户，请修正以下行后重试", "errors": errors}
        )
    db.rollback()  # Release read transaction before expensive password hashing.
    users = [
        User(**fields, password_hash=pwd_context.hash(data.password))
        for _, data, fields in prepared
    ]
    db.execute(text("BEGIN IMMEDIATE"))
    try:
        actor = (
            db.query(User)
            .filter(
                User.id == actor_id,
                User.is_active == 1,
                User.role == "super_admin",
                User.token_version == actor_version,
            )
            .first()
        )
        if not actor:
            raise HTTPException(403, "管理员登录或权限已失效，未导入任何用户")
        usernames = [u.username for u in users]
        if db.query(User.id).filter(User.username.in_(usernames)).first():
            raise HTTPException(409, "账号已被其他操作创建，请刷新后检查表格；未导入任何用户")
        departments = {
            fields["department_id"]
            for _, _, fields in prepared
            if fields["department_id"] is not None
        }
        from models import Department

        current_departments = (
            {
                d.id: d.name
                for d in db.query(Department).filter(Department.id.in_(departments)).all()
            }
            if departments
            else {}
        )
        if len(current_departments) != len(departments):
            raise HTTPException(409, "部门已被其他操作删除，请刷新后检查表格；未导入任何用户")
        for user in users:
            if user.department_id is not None:
                user.department = current_departments[user.department_id]
        db.add_all(users)
        db.add(
            OperationLog(
                user=actor.name,
                module="用户权限",
                action="导入用户",
                detail=f"新增{len(users)}个用户",
            )
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"message": f"成功新增{len(users)}个用户", "count": len(users)}
