from datetime import date
from io import BytesIO
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from deps import require_perm, get_user_permissions, scope_query
from routers.import_export import MODULES

router = APIRouter(prefix='/api/reports', tags=['reports'])


@router.get('/export')
def export_report(db: Session = Depends(get_db), user=Depends(require_perm('reports'))):
    permissions = get_user_permissions(user, db)
    if 'export' not in permissions:
        raise HTTPException(403, '没有导出权限')
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    modules = ['it', 'phone', 'medical', 'number']
    if 'wechat' in permissions and (user.role == 'super_admin' or user.data_scope == 'all'):
        modules.append('wechat')
    distribution = {}
    for module in modules:
        cfg = MODULES[module]
        model = cfg['model']
        rows = scope_query(db.query(model), model, user).all() if hasattr(model, 'department_id') else db.query(model).all()
        fields = [field for field in cfg['fields'] if field[1] != 'wx_password']
        ws = wb.create_sheet(cfg['label'])
        ws.append([field[0] for field in fields])
        for row in rows:
            values = [getattr(row, field[1], None) for field in fields]
            ws.append(values)
            # Keep user input as text, never Excel formulas.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            if module != 'wechat':
                key = row.department or '未分配部门'
                distribution.setdefault(key, dict.fromkeys(['it', 'phone', 'medical', 'number'], 0))[module] += 1
    ws = wb.create_sheet('部门统计')
    ws.append(['部门', 'IT设备', '手机设备', '医疗设备', '电话号码'])
    for department, counts in sorted(distribution.items()):
        ws.append([department] + list(counts.values()))
        ws.cell(ws.max_row, 1).data_type = 's'
    data = BytesIO()
    wb.save(data)
    data.seek(0)
    filename = f'资产报表_{date.today()}.xlsx'
    return StreamingResponse(data, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(filename)}"})
