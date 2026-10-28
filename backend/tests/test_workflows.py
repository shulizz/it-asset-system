"""Run with python -m unittest discover -s backend/tests -v. Uses temporary synthetic data only."""

import json
import os
import sys
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from threading import Event

workspace = tempfile.TemporaryDirectory(prefix="asset-regression-")
os.environ["IT_ASSET_DATABASE_URL"] = "sqlite:///" + str(Path(workspace.name) / "test.db").replace(
    "\\", "/"
)
os.environ["IT_ASSET_SECRET_KEY"] = "isolated-test-secret-never-use-in-production"
os.environ["IT_ASSET_BACKUP_DIR"] = str(Path(workspace.name) / "backups")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from unittest.mock import patch

import openpyxl
from credential_crypto import decrypt_credential, encrypt_credential
from database import Base, SessionLocal, engine
from fastapi.testclient import TestClient
from main import app
from migrations import migrate_workflow
from models import (
    DeleteRequest,
    Department,
    ITAsset,
    MedicalAsset,
    OperationLog,
    Role,
    ScrapRequest,
    TransferRecord,
    User,
    WeChatAccount,
)
from routers.auth import create_token, pwd_context
from sqlalchemy import create_engine, text


class WorkflowTests(unittest.TestCase):
    def test_department_rename_preserves_medical_access_and_report_group(self):
        self.use_medical_departments()
        with SessionLocal() as db:
            db.add(
                ITAsset(
                    asset_number="RENAME-HOSP",
                    name="主机",
                    status="in_use",
                    department_id=1,
                    department="护理部",
                )
            )
            db.commit()
        result = self.client.put(
            "/api/departments/1", headers=self.headers(1), json={"name": "护理中心"}
        )
        self.assertEqual(result.status_code, 200, result.text)
        identity = self.client.get("/api/auth/me", headers=self.headers(2)).json()
        self.assertTrue(identity["medical_access"])
        self.assertEqual(identity["department"], "护理中心")
        self.assertEqual(
            self.client.get("/api/assets/medical", headers=self.headers(2)).status_code, 200
        )
        summary = self.client.get("/api/reports/group-summary", headers=self.headers(1)).json()
        self.assertEqual(summary["groups"][0]["counts"]["it"], 1)
        self.assertIn("护理中心", summary["hospital_departments"])

    def test_all_asset_edits_invalidate_pending_delete_approval(self):
        from models import FireAsset, PhoneAsset, PhoneNumber

        examples = [
            (ITAsset, {"asset_number": "REV-IT", "name": "主机"}),
            (PhoneAsset, {"asset_number": "REV-PH", "brand_model": "手机"}),
            (MedicalAsset, {"asset_number": "REV-MED", "name": "医疗"}),
            (FireAsset, {"asset_number": "REV-FIRE", "name": "灭火器"}),
            (PhoneNumber, {"number": "13800000002"}),
        ]
        with SessionLocal() as db:
            db.get(User, 4).data_scope = "all"
            db.get(User, 4).role = "super_admin"
            db.commit()
        for model, values in examples:
            with SessionLocal() as db:
                item = model(**values, department="甲部门", department_id=1, status="in_use")
                db.add(item)
                db.commit()
                identity = item.id
            requested = self.client.post(
                "/api/delete-request",
                headers=self.headers(1),
                json={
                    "table_name": model.__tablename__,
                    "record_id": identity,
                    "record_desc": "test",
                    "reason": "test",
                },
            )
            self.assertEqual(requested.status_code, 200, requested.text)
            with SessionLocal() as db:
                db.get(model, identity).notes = "申请后修改"
                db.commit()
            response = self.client.put(
                f"/api/delete-request/{requested.json()['id']}/approve", headers=self.headers(4)
            )
            self.assertEqual(response.status_code, 409, response.text)
            with SessionLocal() as db:
                self.assertEqual(db.get(model, identity).status, "in_use")

    def test_notes_edit_preserves_in_use_without_named_person(self):
        identity = self.asset(status="in_use")
        response = self.client.put(
            f"/api/assets/it/{identity}",
            headers=self.headers(3),
            json={"name": "测试设备", "notes": "仅修改备注"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["status"], "in_use")

    def test_scrapped_number_cannot_be_restored_by_edit(self):
        from models import PhoneNumber

        with SessionLocal() as db:
            db.add(
                PhoneNumber(
                    id=1,
                    number="13800000001",
                    department_id=1,
                    department="甲部门",
                    status="scrapped",
                )
            )
            db.commit()
        response = self.client.put(
            "/api/assets/numbers/1",
            headers=self.headers(3),
            json={"number": "13800000001", "notes": "备注"},
        )
        self.assertEqual(response.status_code, 400, response.text)
        with SessionLocal() as db:
            self.assertEqual(db.get(PhoneNumber, 1).status, "scrapped")

    def test_it_scrap_rejects_stale_owner_and_legacy_missing_revision(self):
        for legacy in (False, True):
            identity = self.asset(status="in_use", number=f"STALE-{legacy}")
            request = self.client.post(
                "/api/scrap",
                headers=self.headers(2),
                json={"asset_type": "it", "asset_id": identity, "asset_desc": "test"},
            )
            self.assertEqual(request.status_code, 200, request.text)
            if legacy:
                with SessionLocal() as db:
                    db.get(ScrapRequest, request.json()["id"]).source_revision = None
                    db.commit()
            else:
                edited = self.client.put(
                    f"/api/assets/it/{identity}",
                    headers=self.headers(3),
                    json={"name": "测试设备", "user_name": "新使用人"},
                )
                self.assertEqual(edited.status_code, 200, edited.text)
            approved = self.client.put(
                f"/api/scrap/{request.json()['id']}/approve", headers=self.headers(3)
            )
            self.assertEqual(approved.status_code, 409, approved.text)
            with SessionLocal() as db:
                self.assertEqual(db.get(ITAsset, identity).status, "in_use")

    def test_medical_history_delete_is_hidden_and_cannot_be_rejected(self):
        with SessionLocal() as db:
            db.get(User, 3).data_scope = "all"
            db.add(
                TransferRecord(
                    id=1,
                    transfer_number="MED-TR",
                    type="checkout",
                    asset_desc="medical confidential",
                    asset_type="medical",
                    department_id=2,
                )
            )
            db.add(
                ScrapRequest(
                    id=1,
                    request_number="MED-SC",
                    asset_desc="medical confidential",
                    asset_type="medical",
                    department_id=2,
                )
            )
            db.commit()
        for table in ("transfer_records", "scrap_requests"):
            response = self.client.post(
                "/api/delete-request",
                headers=self.headers(1),
                json={
                    "table_name": table,
                    "record_id": 1,
                    "record_desc": "medical confidential",
                    "reason": "test",
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["asset_type"], "medical")
            identity = response.json()["id"]
            # Old requests with missing type must also remain protected.
            with SessionLocal() as db:
                db.get(DeleteRequest, identity).asset_type = None
                db.commit()
            listing = self.client.get("/api/delete-request", headers=self.headers(3))
            self.assertNotIn("medical confidential", listing.text)
            for action in ("reject", "approve"):
                result = self.client.put(
                    f"/api/delete-request/{identity}/{action}", headers=self.headers(3)
                )
                self.assertEqual(result.status_code, 403, result.text)

    def test_wechat_deletion_requires_independent_approval_and_retains_archive(self):
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.permissions = json.dumps(["wechat", "approval"])
            user.data_scope = "all"
            db.add(WeChatAccount(id=1, wx_account="synthetic-delete", status="idle"))
            db.commit()
        request = self.client.delete("/api/wechat/1", headers=self.headers(3))
        self.assertEqual(request.status_code, 200, request.text)
        with SessionLocal() as db:
            self.assertIsNotNone(db.get(WeChatAccount, 1))
            pending = db.query(DeleteRequest).filter_by(table_name="wechat_accounts").one()
            identity = pending.id
        own = self.client.put(f"/api/delete-request/{identity}/approve", headers=self.headers(3))
        self.assertEqual(own.status_code, 400, own.text)
        approved = self.client.put(
            f"/api/delete-request/{identity}/approve", headers=self.headers(1)
        )
        self.assertEqual(approved.status_code, 200, approved.text)
        with SessionLocal() as db:
            self.assertEqual(db.get(WeChatAccount, 1).status, "archived")
        self.assertEqual(self.client.get("/api/wechat", headers=self.headers(3)).json(), [])

    def test_delegated_fire_manager_has_only_fire_access(self):
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.permissions = json.dumps(["fire"])
            user.data_scope = "all"
            user.role = "消防管理员"
            user.email = "fire-manager@example.com"
            db.commit()
        headers = self.headers(3)
        created = self.client.post(
            "/api/assets/fire",
            headers=headers,
            json={"name": "灭火器", "department": "乙部门", "location": "二楼"},
        )
        self.assertEqual(created.status_code, 200, created.text)
        identity = created.json()["id"]
        self.assertEqual(
            self.client.put(
                f"/api/assets/fire/{identity}",
                headers=headers,
                json={"name": "灭火器", "location": "三楼"},
            ).status_code,
            200,
        )
        for path in [
            "/api/assets/fire",
            "/api/ie/fire/template",
            "/api/ie/fire/export",
            "/api/assets/fire/notice-recipients",
        ]:
            self.assertEqual(self.client.get(path, headers=headers).status_code, 200, path)
        recipients = self.client.get("/api/assets/fire/notice-recipients", headers=headers).json()[
            "users"
        ]
        self.assertIn(3, [user["id"] for user in recipients])
        with patch(
            "notifier.send_fire_expiry_notice",
            return_value={"success": True, "message": "测试", "count": 0},
        ) as send:
            result = self.client.post(
                "/api/assets/fire/send-expiry-email",
                headers=headers,
                json={"recipient_user_ids": [3]},
            )
            self.assertEqual(result.status_code, 200, result.text)
            self.assertEqual(send.call_args.args[2], ["fire-manager@example.com"])
        request = self.client.post(
            "/api/delete-request",
            headers=headers,
            json={
                "table_name": "fire_assets",
                "record_id": identity,
                "record_desc": "灭火器",
                "reason": "测试",
            },
        )
        self.assertEqual(request.status_code, 200, request.text)
        self.assertEqual(
            self.client.put(
                f"/api/delete-request/{request.json()['id']}/approve", headers=headers
            ).status_code,
            403,
        )
        for path in [
            "/api/assets/it",
            "/api/assets/phone",
            "/api/assets/medical",
            "/api/assets/numbers",
            "/api/wechat",
            "/api/reports/group-summary",
            "/api/auth/users",
            "/api/transfer",
            "/api/ie/it/export",
        ]:
            self.assertEqual(self.client.get(path, headers=headers).status_code, 403, path)
        with SessionLocal() as db:
            db.get(User, 3).permissions = "[]"
            db.commit()
        self.assertEqual(self.client.get("/api/assets/fire", headers=headers).status_code, 403)

    def test_scrapped_hosts_monitors_are_separate_from_report_totals(self):
        from models import PhoneAsset

        with SessionLocal() as db:
            db.add(Department(id=3, name="医院"))
            for department, identity, hosts, monitors in [("甲部门", 1, 27, 18), ("医院", 3, 1, 1)]:
                for name, category, count in [
                    ("主机", "台式机", hosts),
                    ("显示器", "显示器", monitors),
                ]:
                    for index in range(count):
                        db.add(
                            ITAsset(
                                asset_number=f"HIST-{identity}-{category}-{index}",
                                name=name,
                                category=category,
                                department_id=identity,
                                department=department,
                                status="scrapped",
                            )
                        )
                db.add(
                    ITAsset(
                        asset_number=f"ACTIVE-{identity}",
                        name="在用主机",
                        department_id=identity,
                        department=department,
                        status="in_use",
                    )
                )
            db.add(
                PhoneAsset(
                    asset_number="PH-SCRAP",
                    brand_model="手机",
                    quantity=5,
                    department_id=1,
                    department="甲部门",
                    status="scrapped",
                )
            )
            db.add(
                PhoneAsset(
                    asset_number="PH-ACTIVE",
                    brand_model="手机",
                    quantity=20,
                    department_id=1,
                    department="甲部门",
                    status="in_use",
                )
            )
            db.commit()
        result = self.client.get("/api/reports/group-summary", headers=self.headers(1))
        self.assertEqual(result.status_code, 200, result.text)
        summary = result.json()
        hospital, network = summary["groups"]
        self.assertEqual((hospital["counts"]["it"], network["counts"]["it"]), (1, 1))
        self.assertEqual(hospital["it_scrapped"], {"hosts": 1, "monitors": 1, "other": 0})
        self.assertEqual(network["it_scrapped"], {"hosts": 27, "monitors": 18, "other": 0})
        self.assertEqual(network["scrapped"]["phone"], 5)
        self.assertEqual(network["counts"]["phone"], 20)
        self.assertEqual(summary["it_scrapped_totals"], {"hosts": 28, "monitors": 19, "other": 0})
        exported = self.client.get("/api/reports/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(exported.content))
        rows = list(book["医院网络部统计"].values)
        self.assertEqual(rows[1][rows[0].index("IT设备")], 1)
        self.assertEqual(rows[2][rows[0].index("手机设备")], 20)
        rows = list(book["部门统计"].values)
        self.assertEqual(sum(row[rows[0].index("IT设备")] for row in rows[1:]), 2)
        rows = list(book["报废分类统计"].values)
        self.assertEqual((rows[1][1], rows[1][2], rows[2][1], rows[2][2]), (1, 1, 27, 18))
        self.assertEqual(
            len(list(book["IT设备"].values)) - 1, 49
        )  # Historical detail remains exportable.

    def test_hospital_network_report_groups_quantities_scope_and_export(self):
        from models import FireAsset, PhoneAsset, PhoneNumber

        with SessionLocal() as db:
            for identity, name in enumerate(["医院", "护理部", "门诊", "办公室"], start=3):
                db.add(Department(id=identity, name=name))
                db.add(
                    ITAsset(
                        asset_number=f"HOSP-{identity}",
                        name="医院设备",
                        department_id=identity,
                        department=name,
                        status="in_use",
                    )
                )
            db.add_all(
                [
                    ITAsset(
                        asset_number="NETWORK",
                        name="网络设备",
                        department_id=1,
                        department="甲部门",
                        status="idle",
                    ),
                    ITAsset(
                        asset_number="BY-ID",
                        name="以部门ID为准",
                        department_id=2,
                        department="医院",
                        status="in_use",
                    ),
                    ITAsset(asset_number="UNASSIGNED", name="未分配", status="in_use"),
                    ITAsset(
                        asset_number="ARCHIVED",
                        name="移出档案",
                        department_id=3,
                        department="医院",
                        status="archived",
                    ),
                    PhoneAsset(
                        asset_number="PH-H",
                        brand_model="手机",
                        department_id=3,
                        department="医院",
                        quantity=12,
                        status="in_use",
                    ),
                    PhoneAsset(
                        asset_number="PH-N",
                        brand_model="手机",
                        department_id=1,
                        department="甲部门",
                        quantity=20,
                        status="idle",
                    ),
                    PhoneNumber(
                        number="10086", department_id=5, department="门诊", status="in_use"
                    ),
                    MedicalAsset(
                        asset_number="MED-H",
                        name="医疗设备",
                        department_id=3,
                        department="医院",
                        status="in_use",
                    ),
                    FireAsset(
                        asset_number="FIRE-H",
                        name="消防设备",
                        department_id=3,
                        department="医院",
                        status="in_use",
                    ),
                    WeChatAccount(wx_account="global-wechat"),
                ]
            )
            db.commit()
        result = self.client.get("/api/reports/group-summary", headers=self.headers(1))
        self.assertEqual(result.status_code, 200, result.text)
        summary = result.json()
        hospital, network = summary["groups"]
        self.assertEqual(
            (hospital["counts"]["it"], network["counts"]["it"], summary["unassigned"]["it"]),
            (4, 2, 1),
        )
        self.assertEqual((hospital["counts"]["phone"], network["counts"]["phone"]), (12, 20))
        self.assertEqual(hospital["counts"]["number"], 1)
        self.assertEqual(hospital["counts"]["medical"], 1)
        self.assertEqual(hospital["counts"]["fire"], 1)
        self.assertEqual(summary["wechat_count"], 1)
        exported = self.client.get("/api/reports/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(exported.content))
        rows = list(book["医院网络部统计"].values)
        self.assertEqual(rows[1][rows[0].index("手机设备")], 12)
        self.assertEqual(rows[2][rows[0].index("手机设备")], 20)
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.permissions = json.dumps(["reports", "export"])
            user.data_scope = "all_read_dept_write"
            db.commit()
        summary = self.client.get("/api/reports/group-summary", headers=self.headers(3)).json()
        self.assertNotIn("medical", summary["groups"][0]["counts"])
        self.assertNotIn("fire", summary["groups"][0]["counts"])
        self.assertIsNone(summary["wechat_count"])
        self.assertEqual(summary["groups"][0]["counts"]["it"], 4)
        with SessionLocal() as db:
            db.get(User, 3).data_scope = "department"
            db.commit()
        summary = self.client.get("/api/reports/group-summary", headers=self.headers(3)).json()
        self.assertEqual(summary["scope"], "department")
        self.assertEqual(summary["groups"][0]["counts"]["it"], 0)
        self.assertEqual(summary["groups"][1]["counts"]["it"], 1)
        self.assertEqual(summary["groups"][1]["counts"]["phone"], 20)
        self.assertEqual(
            self.client.get("/api/reports/group-summary", headers=self.headers(2)).status_code, 403
        )
        self.assertIn(self.client.get("/api/reports/group-summary").status_code, (401, 403))

    def test_fire_location_round_trip_import_export_and_legacy_update(self):
        item = self.create_fire_fixture()
        response = self.client.put(
            f"/api/assets/fire/{item['id']}",
            headers=self.headers(1),
            json={"name": item["name"], "location": "二楼消防通道"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        response = self.client.put(
            f"/api/assets/fire/{item['id']}", headers=self.headers(1), json={"name": "更新名称"}
        )
        self.assertEqual(response.json()["location"], "二楼消防通道")
        exported = self.client.get("/api/ie/fire/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(exported.content))
        rows = list(book.active.values)
        self.assertEqual(rows[1][rows[0].index("位置")], "二楼消防通道")
        wb = openpyxl.Workbook()
        wb.active.append(["消防设备", "位置"])
        wb.active.append(["消防栓", "三楼东侧"])
        buf = BytesIO()
        wb.save(buf)
        response = self.client.post(
            "/api/ie/fire/import",
            headers=self.headers(1),
            files={"file": ("fire.xlsx", buf.getvalue())},
        )
        self.assertEqual(response.json()["imported"], 1, response.text)
        records = self.client.get("/api/assets/fire", headers=self.headers(1)).json()
        self.assertIn("三楼东侧", [row["location"] for row in records])

    def test_fire_email_authorization_selection_dates_and_escaping(self):
        from models import FireAsset

        item = self.create_fire_fixture()
        with SessionLocal() as db:
            db.get(User, 1).email = "admin@example.com"
            db.get(User, 3).email = "ordinary@example.com"
            asset = db.get(FireAsset, item["id"])
            asset.expiry_date = date.today() + timedelta(days=10)
            asset.location = '<img src=x onerror="test">'
            db.add(
                FireAsset(
                    asset_number="NOT-DUE",
                    name="未到期",
                    status="in_use",
                    expiry_date=date.today() + timedelta(days=31),
                )
            )
            db.add(
                FireAsset(
                    asset_number="SCRAPPED-FIRE",
                    name="已报废",
                    status="scrapped",
                    expiry_date=date.today() - timedelta(days=1),
                )
            )
            db.commit()
        with patch("notifier.smtp_configured", return_value=True):
            recipients = self.client.get(
                "/api/assets/fire/notice-recipients", headers=self.headers(1)
            )
            self.assertEqual([u["id"] for u in recipients.json()["users"]], [1])
        with patch(
            "notifier.send_email", return_value={"success": True, "message": "测试已发送"}
        ) as send:
            payload = {"recipient_user_ids": [1, 1]}
            self.assertEqual(
                self.client.post(
                    "/api/assets/fire/send-expiry-email", headers=self.headers(3), json=payload
                ).status_code,
                403,
            )
            self.assertEqual(
                self.client.get(
                    "/api/assets/fire/notice-recipients", headers=self.headers(3)
                ).status_code,
                403,
            )
            self.assertEqual(
                self.client.post(
                    "/api/assets/fire/send-expiry-email",
                    headers=self.headers(1),
                    json={"recipient_user_ids": [3]},
                ).status_code,
                403,
            )
            self.assertEqual(
                self.client.post(
                    "/api/assets/fire/send-expiry-email",
                    headers=self.headers(1),
                    json={"recipient_user_ids": [999]},
                ).status_code,
                400,
            )
            self.assertEqual(
                self.client.post(
                    "/api/assets/fire/send-expiry-email?days=0",
                    headers=self.headers(1),
                    json=payload,
                ).status_code,
                400,
            )
            send.assert_not_called()
            response = self.client.post(
                "/api/assets/fire/send-expiry-email?days=30", headers=self.headers(1), json=payload
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["count"], 1)
            args = send.call_args.args
            self.assertEqual(args[0], ["admin@example.com"])
            self.assertIn("消防设备到期提醒", args[1])
            self.assertIn("&lt;img", args[2])
            self.assertNotIn("<img", args[2])
            self.assertIn("位置", args[2])
            self.assertNotIn("NOT-DUE", args[2])
            self.assertNotIn("SCRAPPED-FIRE", args[2])
        with SessionLocal() as db:
            db.get(FireAsset, item["id"]).expiry_date = date.today() + timedelta(days=60)
            db.commit()
        with patch("notifier.send_email") as send:
            response = self.client.post(
                "/api/assets/fire/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": [1]},
            )
            self.assertEqual(response.json()["count"], 0)
            send.assert_not_called()

    def create_fire_fixture(self):
        response = self.client.post(
            "/api/assets/fire",
            headers=self.headers(1),
            json={
                "name": "干粉灭火器",
                "department": "甲部门",
                "purchase_date": "2026-01-01",
                "expiry_date": "2028-01-01",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_fire_admin_dates_export_and_department_integrity(self):
        item = self.create_fire_fixture()
        self.assertTrue(item["asset_number"].startswith("FIRE-"))
        self.assertEqual(item["status"], "in_use")
        self.assertEqual(item["purchase_date"], "2026-01-01")
        response = self.client.put(
            f"/api/assets/fire/{item['id']}",
            headers=self.headers(1),
            json={"name": "更新灭火器", "purchase_date": "2029-01-01"},
        )
        self.assertEqual(response.status_code, 400, response.text)
        invalid = self.client.post(
            "/api/assets/fire",
            headers=self.headers(1),
            json={"name": "灭火器", "purchase_date": "2026-02-01", "expiry_date": "2025-01-01"},
        )
        self.assertEqual(invalid.status_code, 422)
        response = self.client.get("/api/ie/fire/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(response.content))
        self.assertTrue({"消防设备", "购买时间", "到期时间"}.issubset(next(book.active.values)))
        report = self.client.get("/api/reports/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(report.content))
        self.assertIn("消防设备", book.sheetnames)
        renamed = self.client.put(
            "/api/departments/1", headers=self.headers(1), json={"name": "新部门"}
        )
        self.assertEqual(renamed.status_code, 200)
        row = self.client.get("/api/assets/fire", headers=self.headers(1)).json()[0]
        self.assertEqual(row["department"], "新部门")

    def test_fire_ordinary_user_cannot_access_even_with_broad_existing_grants(self):
        item = self.create_fire_fixture()
        from permissions import PERMISSION_KEYS

        with SessionLocal() as db:
            user = db.get(User, 3)
            user.permissions = json.dumps(sorted(PERMISSION_KEYS - {"fire"}))
            user.role = "管理员"
            user.data_scope = "all"
            db.commit()
        headers = self.headers(3)
        for path in ["/api/assets/fire", "/api/ie/fire/template", "/api/ie/fire/export"]:
            self.assertEqual(self.client.get(path, headers=headers).status_code, 403, path)
        self.assertEqual(
            self.client.post(
                "/api/assets/fire", headers=headers, json={"name": "禁止新增"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(
                f"/api/assets/fire/{item['id']}", headers=headers, json={"name": "禁止修改"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/scrap",
                headers=headers,
                json={"asset_type": "fire", "asset_id": item["id"], "asset_desc": "禁止报废"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/delete-request",
                headers=headers,
                json={
                    "table_name": "fire_assets",
                    "record_id": item["id"],
                    "record_desc": "禁止删除",
                    "reason": "test",
                },
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/transfer",
                headers=headers,
                json={
                    "asset_type": "fire",
                    "asset_id": item["id"],
                    "asset_desc": "禁止调拨",
                    "transfer_number": "x",
                    "type": "transfer",
                    "new_user": "某人",
                },
            ).status_code,
            403,
        )
        report = self.client.get("/api/reports/export", headers=headers)
        book = openpyxl.load_workbook(BytesIO(report.content))
        self.assertNotIn("消防设备", book.sheetnames)
        self.assertNotIn("消防设备", next(book["部门统计"].values))
        self.assertEqual(self.client.get("/api/logs", headers=headers).json(), [])
        self.assertIn(self.client.get("/api/assets/fire").status_code, (401, 403))

    def test_fire_import_and_hidden_workflow_records(self):
        wb = openpyxl.Workbook()
        wb.active.append(["消防设备", "购买时间", "到期时间", "部门"])
        wb.active.append(["灭火器", "2026-01-01", "2028-01-01", "甲部门"])
        wb.active.append(["错误日期", "2026-01-01", "2025-01-01", "甲部门"])
        buf = BytesIO()
        wb.save(buf)
        upload = {
            "file": (
                "fire.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        self.assertEqual(
            self.client.post(
                "/api/ie/fire/import", headers=self.headers(3), files=upload
            ).status_code,
            403,
        )
        result = self.client.post("/api/ie/fire/import", headers=self.headers(1), files=upload)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual((result.json()["imported"], result.json()["skipped"]), (1, 1))
        asset = self.client.get("/api/assets/fire", headers=self.headers(1)).json()[0]
        scrap = self.client.post(
            "/api/scrap",
            headers=self.headers(1),
            json={"asset_type": "fire", "asset_id": asset["id"], "asset_desc": "test"},
        )
        self.assertEqual(scrap.status_code, 200, scrap.text)
        self.assertEqual(self.client.get("/api/scrap", headers=self.headers(3)).json(), [])
        self.assertEqual(
            self.client.put(
                f"/api/scrap/{scrap.json()['id']}/reject", headers=self.headers(3)
            ).status_code,
            403,
        )
        delete = self.client.post(
            "/api/delete-request",
            headers=self.headers(1),
            json={
                "table_name": "fire_assets",
                "record_id": asset["id"],
                "record_desc": "test",
                "reason": "test",
            },
        )
        self.assertEqual(delete.status_code, 200, delete.text)
        self.assertEqual(self.client.get("/api/delete-request", headers=self.headers(3)).json(), [])
        self.assertEqual(
            self.client.put(
                f"/api/delete-request/{delete.json()['id']}/approve", headers=self.headers(3)
            ).status_code,
            403,
        )

    def test_phone_quantity_additive_migration_preserves_legacy_rows(self):
        from migrations import migrate_schema

        with tempfile.TemporaryDirectory() as folder:
            old_engine = create_engine("sqlite:///" + str(Path(folder) / "legacy.db"))
            Base.metadata.create_all(old_engine)
            with old_engine.begin() as connection:
                connection.execute(text("ALTER TABLE phone_assets DROP COLUMN quantity"))
                connection.execute(text("ALTER TABLE phone_assets DROP COLUMN revision"))
                connection.execute(
                    text(
                        "INSERT INTO phone_assets (asset_number,brand_model,status) VALUES ('LEGACY','原手机','idle')"
                    )
                )
                connection.execute(
                    text(
                        "INSERT INTO it_assets (asset_number,name,department) VALUES ('EMPTY','空部门设备','')"
                    )
                )
            migrate_schema(old_engine)
            migrate_schema(old_engine)
            with old_engine.connect() as connection:
                row = connection.execute(
                    text("SELECT asset_number,brand_model,quantity,revision FROM phone_assets")
                ).one()
                self.assertEqual(tuple(row), ("LEGACY", "原手机", 1, 0))
                self.assertEqual(
                    connection.execute(
                        text("SELECT department FROM it_assets WHERE asset_number='EMPTY'")
                    ).scalar_one(),
                    "",
                )
            old_engine.dispose()

    def test_phone_approved_application_splits_quantity(self):
        phone = self.create_phone_batch()
        request = self.client.post(
            "/api/delete-request",
            headers=self.headers(2),
            json={
                "table_name": "apply_requests",
                "record_desc": "手机调拨",
                "reason": "测试",
                "application_type": "transfer",
                "asset_type": "phone",
                "asset_id": phone["id"],
                "target_user": "李四",
                "target_department_id": 1,
                "quantity": 4,
            },
        )
        self.assertEqual(request.status_code, 200, request.text)
        result = self.client.put(
            f"/api/delete-request/{request.json()['id']}/approve", headers=self.headers(3)
        )
        self.assertEqual(result.status_code, 200, result.text)
        from models import PhoneAsset

        with SessionLocal() as db:
            self.assertEqual(db.get(PhoneAsset, phone["id"]).quantity, 16)
            target = db.query(PhoneAsset).filter(PhoneAsset.user_name == "李四").one()
            self.assertEqual(target.quantity, 4)

    def create_phone_batch(self, quantity=20):
        response = self.client.post(
            "/api/assets/phone",
            headers=self.headers(3),
            json={
                "brand_model": "同型号测试手机",
                "user_name": "张三",
                "department": "甲部门",
                "quantity": quantity,
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_phone_partial_scrap_conserves_quantity_and_print_request(self):
        phone = self.create_phone_batch()
        response = self.client.post(
            "/api/scrap",
            headers=self.headers(2),
            json={
                "asset_type": "phone",
                "asset_id": phone["id"],
                "asset_desc": "test",
                "quantity": 3,
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        request = response.json()
        self.assertEqual(request["quantity"], 3)
        self.assertEqual(
            self.client.put(
                f"/api/scrap/{request['id']}/approve", headers=self.headers(3)
            ).status_code,
            200,
        )
        from models import PhoneAsset

        with SessionLocal() as db:
            source = db.get(PhoneAsset, phone["id"])
            self.assertEqual(
                (source.quantity, source.status, source.user_name), (17, "in_use", "张三")
            )
            scraps = db.query(PhoneAsset).filter(PhoneAsset.status == "scrapped").all()
            self.assertEqual([(x.quantity, x.department_id) for x in scraps], [(3, 1)])
            record = db.query(TransferRecord).one()
            self.assertEqual(
                (record.quantity, record.asset_id, record.result_asset_id),
                (3, source.id, scraps[0].id),
            )
            self.assertEqual(sum(x.quantity for x in db.query(PhoneAsset).all()), 20)
        self.assertEqual(
            self.client.put(
                f"/api/scrap/{request['id']}/approve", headers=self.headers(3)
            ).status_code,
            400,
        )

    def test_phone_partial_transfer_return_and_full_transition(self):
        phone = self.create_phone_batch()

        def transfer(asset_id, quantity, action, recipient="李四", department="乙部门"):
            return self.client.post(
                "/api/transfer",
                headers=self.headers(1),
                json={
                    "transfer_number": "ignored",
                    "type": action,
                    "asset_desc": "ignored",
                    "asset_type": "phone",
                    "asset_id": asset_id,
                    "quantity": quantity,
                    "new_user": recipient,
                    "new_dept": department,
                },
            )

        moved = transfer(phone["id"], 5, "transfer")
        self.assertEqual(moved.status_code, 200, moved.text)
        target_id = moved.json()["result_asset_id"]
        self.assertNotEqual(target_id, phone["id"])
        from models import PhoneAsset

        with SessionLocal() as db:
            self.assertEqual(
                (
                    db.get(PhoneAsset, phone["id"]).quantity,
                    db.get(PhoneAsset, phone["id"]).department_id,
                ),
                (15, 1),
            )
            target = db.get(PhoneAsset, target_id)
            self.assertEqual(
                (target.quantity, target.department_id, target.user_name), (5, 2, "李四")
            )
        returned = transfer(target_id, 2, "return", "", "")
        self.assertEqual(returned.status_code, 200, returned.text)
        idle_id = returned.json()["result_asset_id"]
        self.assertEqual(transfer(idle_id, 2, "checkout", "王五").status_code, 200)
        with SessionLocal() as db:
            self.assertEqual(sum(x.quantity for x in db.query(PhoneAsset)), 20)
            self.assertEqual(db.get(PhoneAsset, idle_id).user_name, "王五")

    def test_phone_quantity_validation_scope_and_legacy_edit(self):
        for quantity in [0, -1, 1.5, True, "3", 100001]:
            response = self.client.post(
                "/api/assets/phone",
                headers=self.headers(3),
                json={"brand_model": "test", "department": "甲部门", "quantity": quantity},
            )
            self.assertEqual(response.status_code, 422, response.text)
        phone = self.create_phone_batch()
        response = self.client.put(
            f"/api/assets/phone/{phone['id']}",
            headers=self.headers(3),
            json={"brand_model": "更新型号"},
        )
        self.assertEqual(response.json()["quantity"], 20)
        self.assertEqual(
            self.client.post(
                "/api/assets/phone",
                headers=self.headers(3),
                json={"brand_model": "bad", "department": "乙部门", "quantity": 10},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/assets/phone",
                headers=self.headers(3),
                json={
                    "brand_model": "bad",
                    "department": "甲部门",
                    "quantity": 10,
                    "imei": "single",
                },
            ).status_code,
            400,
        )
        payload = {
            "asset_type": "phone",
            "asset_id": phone["id"],
            "asset_desc": "test",
            "quantity": 21,
        }
        self.assertEqual(
            self.client.post("/api/scrap", headers=self.headers(2), json=payload).status_code, 409
        )

    def test_phone_pending_request_rejects_changed_batch(self):
        phone = self.create_phone_batch()
        payload = {
            "asset_type": "phone",
            "asset_id": phone["id"],
            "asset_desc": "test",
            "quantity": 12,
        }
        first = self.client.post("/api/scrap", headers=self.headers(2), json=payload).json()
        second = self.client.post("/api/scrap", headers=self.headers(2), json=payload).json()
        self.assertEqual(
            self.client.put(
                f"/api/scrap/{first['id']}/approve", headers=self.headers(3)
            ).status_code,
            200,
        )
        response = self.client.put(f"/api/scrap/{second['id']}/approve", headers=self.headers(3))
        self.assertEqual(response.status_code, 409, response.text)
        from models import PhoneAsset, ScrapRequest

        with SessionLocal() as db:
            self.assertEqual(db.get(PhoneAsset, phone["id"]).quantity, 8)
            self.assertEqual(db.get(ScrapRequest, second["id"]).status, "pending_approval")

    def test_phone_excel_quantity_round_trip_and_department_totals(self):
        wb = openpyxl.Workbook()
        wb.active.append(["品牌型号", "部门", "使用人", "数量"])
        for value in [20, 0, 2.5, "bad"]:
            wb.active.append(["Excel手机", "甲部门", "张三", value])
        buf = BytesIO()
        wb.save(buf)
        response = self.client.post(
            "/api/ie/phone/import",
            headers=self.headers(1),
            files={
                "file": (
                    "phone.xlsx",
                    buf.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual((response.json()["imported"], response.json()["skipped"]), (1, 3))
        exported = self.client.get("/api/ie/phone/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(exported.content))
        rows = list(book.active.values)
        self.assertEqual(rows[1][rows[0].index("数量")], 20)
        report = self.client.get("/api/reports/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(report.content))
        rows = list(book["部门统计"].values)
        self.assertEqual(rows[1][rows[0].index("手机设备")], 20)

    def test_wechat_basic_edit_is_independent_of_asset_scope_and_secrets(self):
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.data_scope = "all_read_dept_write"
            user.permissions = json.dumps(["wechat_read", "wechat_edit", "assets_write"])
            db.add(
                WeChatAccount(
                    id=1,
                    wx_account="basic-wechat",
                    wx_password=encrypt_credential("private-fixture"),
                    real_name="原实名",
                    purpose="原用途",
                )
            )
            db.commit()
        foreign = self.asset(department_id=2, number="FOREIGN-WECHAT")
        headers = self.headers(3)
        result = self.client.get("/api/wechat", headers=headers)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertIsNone(result.json()[0]["wx_password"])
        updated = self.client.put(
            "/api/wechat/1",
            headers=headers,
            json={"wx_account": "basic-wechat", "purpose": "新用途"},
        )
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertEqual(updated.json()["real_name"], "原实名")
        self.assertIsNone(updated.json()["wx_password"])
        self.assertEqual(
            self.client.post(
                "/api/wechat", headers=headers, json={"wx_account": "new-basic", "purpose": "新增"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                "/api/wechat/1",
                headers=headers,
                json={"wx_account": "basic-wechat", "wx_password": "forbidden-secret"},
            ).status_code,
            403,
        )
        self.assertEqual(self.client.delete("/api/wechat/1", headers=headers).status_code, 403)
        self.assertEqual(
            self.client.put(
                f"/api/assets/it/{foreign}", headers=headers, json={"name": "禁止跨部门"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(
                "/api/wechat/1", headers=headers, json={"wx_account": "new-basic"}
            ).status_code,
            409,
        )
        with SessionLocal() as db:
            self.assertEqual(
                decrypt_credential(db.get(WeChatAccount, 1).wx_password), "private-fixture"
            )
            self.assertEqual(db.get(User, 3).data_scope, "all_read_dept_write")

    def test_wechat_read_only_does_not_allow_edit_and_requires_explicit_grants(self):
        with SessionLocal() as db:
            user = db.get(User, 2)
            user.permissions = json.dumps(["wechat_read"])
            db.add(
                WeChatAccount(
                    id=1,
                    wx_account="readonly-wechat",
                    wx_password=encrypt_credential("private-fixture"),
                )
            )
            db.commit()
        headers = self.headers(2)
        self.assertEqual(self.client.get("/api/wechat", headers=headers).status_code, 200)
        self.assertEqual(
            self.client.post(
                "/api/wechat", headers=headers, json={"wx_account": "denied"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(
                "/api/wechat/1", headers=headers, json={"wx_account": "denied"}
            ).status_code,
            403,
        )
        self.assertEqual(self.client.delete("/api/wechat/1", headers=headers).status_code, 403)
        self.assertEqual(self.client.get("/api/ie/wechat/export", headers=headers).status_code, 403)
        payload = {
            "username": "wechat-editor",
            "name": "编辑",
            "password": "testing-password-123",
            "department": "甲部门",
            "data_scope": "all_read_dept_write",
            "permissions": ["wechat_read", "wechat_edit"],
        }
        self.assertEqual(
            self.client.post("/api/auth/users", headers=self.headers(1), json=payload).status_code,
            200,
        )
        payload.update(username="invalid-wechat-editor", permissions=["wechat_edit"])
        self.assertEqual(
            self.client.post("/api/auth/users", headers=self.headers(1), json=payload).status_code,
            400,
        )

    def test_scrap_application_permission_does_not_grant_edit_approval_or_transfer(self):
        from models import ScrapRequest

        with SessionLocal() as db:
            user = db.get(User, 2)
            user.permissions = json.dumps(["scrap_apply"])
            user.data_scope = "all_read_dept_write"
            db.commit()
        own = self.asset(number="SCRAP-ONLY")
        other = self.asset(department_id=2, number="SCRAP-OTHER")
        headers = self.headers(2)
        self.assertEqual(self.client.get("/api/assets/it", headers=headers).status_code, 200)
        request = self.client.post(
            "/api/scrap",
            headers=headers,
            json={
                "asset_type": "it",
                "asset_id": own,
                "asset_desc": "错误的客户端描述",
                "applicant": "伪造申请人",
                "status": "approved",
                "auditor": "伪造审批人",
            },
        )
        self.assertEqual(request.status_code, 200, request.text)
        self.assertEqual(request.json()["status"], "pending_approval")
        self.assertEqual(request.json()["asset_desc"], "SCRAP-ONLY 测试设备")
        self.assertEqual(request.json()["applicant_id"], 2)
        self.assertIsNone(request.json()["auditor"])
        request_id = request.json()["id"]
        self.assertEqual(
            self.client.put(f"/api/scrap/{request_id}/approve", headers=headers).status_code, 403
        )
        self.assertEqual(
            self.client.put(
                f"/api/assets/it/{own}", headers=headers, json={"name": "不允许编辑"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/transfer",
                headers=headers,
                json={
                    "transfer_number": "forbidden",
                    "type": "checkout",
                    "asset_desc": "设备",
                    "asset_type": "it",
                    "asset_id": own,
                    "new_user": "本人",
                },
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/delete-request",
                headers=headers,
                json={
                    "table_name": "apply_requests",
                    "record_desc": "不能申请领用",
                    "reason": "领用",
                    "application_type": "checkout",
                    "asset_type": "it",
                    "asset_id": own,
                    "target_user": "本人",
                },
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/scrap",
                headers=headers,
                json={"asset_type": "it", "asset_id": other, "asset_desc": "外部门"},
            ).status_code,
            403,
        )
        with SessionLocal() as db:
            self.assertEqual(db.get(ITAsset, own).status, "idle")
            db.add(
                ScrapRequest(
                    request_number="OTHER-PERSON",
                    asset_desc="他人申请",
                    applicant_id=3,
                    asset_type="it",
                    department_id=1,
                )
            )
            db.commit()
        rows = self.client.get("/api/scrap", headers=headers)
        self.assertEqual([r["id"] for r in rows.json()], [request_id])
        approved = self.client.put(f"/api/scrap/{request_id}/approve", headers=self.headers(3))
        self.assertEqual(approved.status_code, 200, approved.text)

    def test_scrap_application_permission_keeps_medical_and_department_guards(self):
        with SessionLocal() as db:
            db.get(User, 2).permissions = json.dumps(["scrap_apply"])
            db.add(
                MedicalAsset(
                    id=1,
                    asset_number="MED-SCRAP",
                    name="医疗设备",
                    department="甲部门",
                    department_id=1,
                    status="idle",
                )
            )
            db.commit()
        payload = {"asset_type": "medical", "asset_id": 1, "asset_desc": "医疗设备"}
        self.assertEqual(
            self.client.get("/api/assets/medical", headers=self.headers(2)).status_code, 403
        )
        self.assertEqual(
            self.client.post("/api/scrap", headers=self.headers(2), json=payload).status_code, 403
        )
        self.use_medical_departments()
        self.assertEqual(
            self.client.get("/api/assets/medical", headers=self.headers(2)).status_code, 200
        )
        self.assertEqual(
            self.client.post("/api/scrap", headers=self.headers(2), json=payload).status_code, 200
        )

    def test_scrap_application_permission_is_configurable_and_importable(self):
        result = self.client.get("/api/auth/permissions", headers=self.headers(1))
        self.assertIn({"key": "scrap_apply", "label": "提交报废申请"}, result.json())
        payload = {
            "username": "scrap-only",
            "name": "申请人",
            "password": "testing-password-123",
            "role": "普通用户",
            "department": "甲部门",
            "data_scope": "all_read_dept_write",
            "permissions": ["scrap_apply"],
        }
        result = self.client.post("/api/auth/users", headers=self.headers(1), json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["permissions"], ["scrap_apply"])
        content = self.user_workbook(
            [
                [
                    "import-scrap",
                    "申请人2",
                    "testing-password-123",
                    "",
                    "普通用户",
                    "甲部门",
                    "仅本部门",
                    "提交报废申请",
                ]
            ]
        )
        result = self.client.post(
            "/api/user-import/import",
            headers=self.headers(1),
            files={"file": ("users.xlsx", content)},
        )
        self.assertEqual(result.status_code, 200, result.text)
        with SessionLocal() as db:
            self.assertEqual(
                json.loads(
                    db.query(User).filter(User.username == "import-scrap").one().permissions
                ),
                ["scrap_apply"],
            )

    def test_self_password_requires_current_password_and_authenticated_identity(self):
        headers = self.headers(2)
        with SessionLocal() as db:
            before = (db.get(User, 2).password_hash, db.get(User, 2).token_version)
        verify = self.client.post(
            "/api/auth/password/verify", headers=headers, json={"current_password": "incorrect"}
        )
        self.assertEqual(verify.status_code, 400)
        self.assertEqual(
            self.client.post(
                "/api/auth/password/change",
                headers=headers,
                json={
                    "current_password": "incorrect",
                    "new_password": "new-password-123",
                    "confirm_password": "new-password-123",
                },
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/password/verify", json={"current_password": "testing-password-123"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/password/change",
                json={
                    "current_password": "testing-password-123",
                    "new_password": "new-password-123",
                    "confirm_password": "new-password-123",
                },
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/password/verify",
                headers=headers,
                json={"current_password": "testing-password-123"},
            ).json(),
            {"verified": True},
        )
        with SessionLocal() as db:
            self.assertEqual((db.get(User, 2).password_hash, db.get(User, 2).token_version), before)

    def test_self_password_change_invalidates_all_old_tokens_and_changes_only_actor(self):
        with SessionLocal() as db:
            user = db.get(User, 2)
            user.permissions = "[]"
            db.commit()
            admin_hash = db.get(User, 1).password_hash
            before = (user.department_id, user.role, user.permissions)
        headers = self.headers(2)
        result = self.client.post(
            "/api/auth/password/change",
            headers=headers,
            json={
                "current_password": "testing-password-123",
                "new_password": "new-password-123",
                "confirm_password": "new-password-123",
                "user_id": 1,
            },
        )
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(self.client.get("/api/auth/me", headers=headers).status_code, 401)
        self.assertEqual(
            self.client.post(
                "/api/auth/login",
                json={"username": "applicant", "password": "testing-password-123"},
            ).status_code,
            401,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "applicant", "password": "new-password-123"}
            ).status_code,
            200,
        )
        with SessionLocal() as db:
            user = db.get(User, 2)
            self.assertTrue(pwd_context.verify("new-password-123", user.password_hash))
            self.assertEqual(db.get(User, 1).password_hash, admin_hash)
            self.assertEqual((user.department_id, user.role, user.permissions), before)
            log = db.query(OperationLog).filter(OperationLog.action == "修改本人密码").one()
            self.assertNotIn("new-password-123", log.detail)
            self.assertNotIn("testing-password-123", log.detail)

    def test_self_password_validation_rejects_mismatch_reuse_and_bcrypt_truncation(self):
        headers = self.headers(2)
        for new, confirm in [
            ("short", "short"),
            ("new-password-123", "different-password"),
            ("testing-password-123", "testing-password-123"),
            ("长" * 25, "长" * 25),
        ]:
            with self.subTest(new_length=len(new)):
                result = self.client.post(
                    "/api/auth/password/change",
                    headers=headers,
                    json={
                        "current_password": "testing-password-123",
                        "new_password": new,
                        "confirm_password": confirm,
                    },
                )
                self.assertEqual(result.status_code, 400, result.text)
                self.assertNotIn(new, result.text)

    def test_self_password_verification_throttles_repeated_guesses(self):
        headers = self.headers(2)
        with patch("routers.auth.pwd_context.verify", return_value=False):
            for _ in range(10):
                self.assertEqual(
                    self.client.post(
                        "/api/auth/password/verify",
                        headers=headers,
                        json={"current_password": "incorrect"},
                    ).status_code,
                    400,
                )
            result = self.client.post(
                "/api/auth/password/change",
                headers=headers,
                json={
                    "current_password": "incorrect",
                    "new_password": "new-password-123",
                    "confirm_password": "new-password-123",
                },
            )
            self.assertEqual(result.status_code, 429)
            self.assertEqual(result.headers["Retry-After"], "60")

    def test_concurrent_admin_reset_is_not_overwritten_by_self_password_change(self):
        entered, release = Event(), Event()
        new_hash = pwd_context.hash("new-password-123")
        reset_hash = pwd_context.hash("admin-reset-123")
        headers = self.headers(2)

        def slow_hash(*args):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Hash wait timed out")
            return new_hash

        with (
            patch("routers.auth.pwd_context.hash", side_effect=slow_hash),
            ThreadPoolExecutor(max_workers=1) as pool,
        ):
            changing = pool.submit(
                self.client.post,
                "/api/auth/password/change",
                headers=headers,
                json={
                    "current_password": "testing-password-123",
                    "new_password": "new-password-123",
                    "confirm_password": "new-password-123",
                },
            )
            try:
                self.assertTrue(entered.wait(3))
                with SessionLocal() as db:
                    db.execute(text("BEGIN IMMEDIATE"))
                    user = db.get(User, 2)
                    user.password_hash = reset_hash
                    user.token_version += 1
                    db.commit()
            finally:
                release.set()
            self.assertEqual(changing.result().status_code, 409)
        with SessionLocal() as db:
            self.assertEqual(db.get(User, 2).password_hash, reset_hash)

    def use_medical_departments(self):
        with SessionLocal() as db:
            for department_id, name in [(1, "护理部"), (2, "医院")]:
                db.get(Department, department_id).name = name
                db.get(Department, department_id).medical_access = 1
                db.get(Department, department_id).report_group = "hospital"
                db.query(User).filter(User.department_id == department_id).update(
                    {User.department: name}
                )
            db.commit()

    def configure_global_reader(self):
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.data_scope = "all_read_dept_write"
            user.permissions = json.dumps(
                [
                    "assets",
                    "assets_write",
                    "reports",
                    "import",
                    "export",
                    "transfer",
                    "approval",
                    "scrap",
                ]
            )
            db.commit()

    def test_medical_visibility_cannot_be_granted_by_global_scope_or_role_label(self):
        from models import PhoneAsset

        self.configure_global_reader()
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.role = "管理员"
            user.data_scope = "all"
            user.permissions = json.dumps(
                [
                    "assets",
                    "assets_write",
                    "reports",
                    "export",
                    "import",
                    "approval",
                    "transfer",
                    "scrap",
                    "logs",
                ]
            )
            db.add(
                MedicalAsset(
                    id=1,
                    asset_number="SECRET-MED",
                    name="受限医疗设备",
                    department="甲部门",
                    department_id=1,
                    status="idle",
                    expiry_date=date.today(),
                )
            )
            db.add(
                PhoneAsset(
                    id=1,
                    asset_number="PUBLIC-PHONE",
                    brand_model="普通手机",
                    department="甲部门",
                    department_id=1,
                    status="idle",
                )
            )
            db.commit()
        headers = self.headers(3)
        for endpoint in [
            "/api/assets/medical",
            "/api/assets/medical/expiring",
            "/api/assets/medical/notice-recipients",
            "/api/ie/medical/export",
            "/api/ie/medical/template",
        ]:
            result = self.client.get(endpoint, headers=headers)
            self.assertEqual(result.status_code, 403, (endpoint, result.text))
        self.assertEqual(
            self.client.post(
                "/api/assets/medical",
                headers=headers,
                json={"name": "禁止新增", "department": "甲部门"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(
                "/api/assets/medical/1", headers=headers, json={"name": "禁止修改"}
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                "/api/scrap",
                headers=headers,
                json={"asset_type": "medical", "asset_id": 1, "asset_desc": "不允许申请"},
            ).status_code,
            403,
        )
        result = self.client.get("/api/reports/export", headers=headers)
        self.assertEqual(result.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(result.content), read_only=True)
        self.assertNotIn("医疗设备", book.sheetnames)
        self.assertNotIn("医疗设备", list(next(book["部门统计"].values)))
        self.assertNotIn("SECRET-MED", str([list(sheet.values) for sheet in book]))
        book.close()
        self.assertEqual(self.client.get("/api/assets/phone", headers=headers).status_code, 200)

    def test_medical_records_and_logs_are_hidden_from_non_medical_departments(self):
        from models import ScrapRequest

        self.configure_global_reader()
        with SessionLocal() as db:
            user = db.get(User, 3)
            user.data_scope = "all"
            user.permissions = json.dumps(["approval", "transfer", "scrap", "logs"])
            db.add_all(
                [
                    ScrapRequest(
                        id=1,
                        request_number="MED-SCRAP",
                        asset_desc="秘密医疗",
                        asset_type="medical",
                        department_id=1,
                    ),
                    ScrapRequest(
                        id=2,
                        request_number="IT-SCRAP",
                        asset_desc="普通电脑",
                        asset_type="it",
                        department_id=1,
                    ),
                    TransferRecord(
                        transfer_number="MED-FLOW",
                        type="checkout",
                        asset_desc="秘密医疗",
                        asset_type="medical",
                        department_id=1,
                    ),
                    DeleteRequest(
                        id=1,
                        table_name="medical_assets",
                        record_id=1,
                        record_desc="秘密医疗",
                        reason="删除",
                        department_id=1,
                        applicant_id=2,
                    ),
                    DeleteRequest(
                        id=2,
                        table_name="apply_requests",
                        record_id=0,
                        asset_type="medical",
                        asset_id=1,
                        record_desc="秘密医疗申请",
                        reason="报废",
                        department_id=1,
                        applicant_id=2,
                    ),
                    OperationLog(
                        user="管理员", module="医疗设备", action="新增", detail="秘密医疗"
                    ),
                    OperationLog(
                        user="管理员", module="设备流转", action="领用", detail="秘密医疗"
                    ),
                    OperationLog(user="管理员", module="IT设备", action="新增", detail="普通电脑"),
                ]
            )
            db.commit()
        headers = self.headers(3)
        for endpoint in ["/api/scrap", "/api/transfer", "/api/delete-request", "/api/logs"]:
            result = self.client.get(endpoint, headers=headers)
            self.assertEqual(result.status_code, 200)
            self.assertNotIn("秘密医疗", result.text)
        self.assertEqual(len(self.client.get("/api/scrap", headers=headers).json()), 1)
        self.assertEqual(
            self.client.put("/api/delete-request/1/reject", headers=headers).status_code, 403
        )
        self.assertEqual(self.client.put("/api/scrap/1/reject", headers=headers).status_code, 403)

    def test_hospital_and_nursing_can_view_medical_but_write_only_own_department(self):
        self.use_medical_departments()
        self.configure_global_reader()
        with SessionLocal() as db:
            db.add_all(
                [
                    MedicalAsset(
                        id=1,
                        asset_number="NURSING-MED",
                        name="护理设备",
                        department="护理部",
                        department_id=1,
                        status="idle",
                    ),
                    MedicalAsset(
                        id=2,
                        asset_number="HOSPITAL-MED",
                        name="医院设备",
                        department="医院",
                        department_id=2,
                        status="idle",
                    ),
                ]
            )
            db.get(User, 4).data_scope = "all_read_dept_write"
            db.get(User, 4).permissions = json.dumps(["assets", "assets_write"])
            db.commit()
        for user_id, own, other in [(3, 1, 2), (4, 2, 1)]:
            headers = self.headers(user_id)
            result = self.client.get("/api/assets/medical", headers=headers)
            self.assertEqual(result.status_code, 200, result.text)
            self.assertEqual(len(result.json()), 2)
            self.assertEqual(
                self.client.put(
                    f"/api/assets/medical/{own}", headers=headers, json={"name": "允许编辑"}
                ).status_code,
                200,
            )
            self.assertEqual(
                self.client.put(
                    f"/api/assets/medical/{other}", headers=headers, json={"name": "禁止编辑"}
                ).status_code,
                403,
            )

    def test_medical_emails_cannot_disclose_to_other_departments(self):
        with SessionLocal() as db:
            db.get(User, 2).email = "other@example.com"
            db.commit()
        recipients = self.client.get(
            "/api/assets/medical/notice-recipients", headers=self.headers(1)
        )
        self.assertEqual(recipients.json()["users"], [])
        with patch("notifier.send_expiry_notice") as send:
            result = self.client.post(
                "/api/assets/medical/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": [2]},
            )
            self.assertEqual(result.status_code, 403)
            send.assert_not_called()

    def test_global_reader_sees_all_but_edits_only_original_department(self):
        self.configure_global_reader()
        own = self.asset(department_id=1, number="OWN-READ")
        other = self.asset(department_id=2, number="OTHER-READ")
        headers = self.headers(3)
        rows = self.client.get("/api/assets/it", headers=headers)
        self.assertEqual(rows.status_code, 200)
        self.assertEqual({row["id"] for row in rows.json()}, {own, other})
        for payload in [{"name": "禁止编辑"}, {"name": "禁止改归属", "department": "甲部门"}]:
            result = self.client.put(f"/api/assets/it/{other}", headers=headers, json=payload)
            self.assertEqual(result.status_code, 403, result.text)
        self.assertEqual(
            self.client.put(
                f"/api/assets/it/{own}", headers=headers, json={"name": "本部门修改"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                f"/api/assets/it/{own}",
                headers=headers,
                json={"name": "本部门修改", "department": "乙部门"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/assets/it",
                headers=headers,
                json={"name": "外部门新增", "department": "乙部门"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/assets/it",
                headers=headers,
                json={"name": "本部门新增", "department": "甲部门"},
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/scrap",
                headers=headers,
                json={"asset_type": "it", "asset_id": other, "asset_desc": "外部门"},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/delete-request",
                headers=headers,
                json={
                    "table_name": "it_assets",
                    "record_id": other,
                    "record_desc": "外部门",
                    "reason": "不允许",
                },
            ).status_code,
            403,
        )
        with SessionLocal() as db:
            self.assertEqual(db.get(ITAsset, other).department_id, 2)
            self.assertNotEqual(db.get(ITAsset, other).name, "禁止编辑")

    def test_global_reader_boundary_applies_to_all_asset_types_and_import(self):
        from models import PhoneAsset, PhoneNumber

        self.configure_global_reader()
        with SessionLocal() as db:
            db.add_all(
                [
                    PhoneAsset(
                        id=1,
                        asset_number="P-FOREIGN",
                        brand_model="其他手机",
                        department="乙部门",
                        department_id=2,
                    ),
                    MedicalAsset(
                        id=1,
                        asset_number="M-FOREIGN",
                        name="其他医疗",
                        department="乙部门",
                        department_id=2,
                    ),
                    PhoneNumber(id=1, number="13800000001", department="乙部门", department_id=2),
                ]
            )
            db.commit()
        headers = self.headers(3)
        for kind, payload in [
            ("phone", {"brand_model": "禁止", "department": "甲部门"}),
            ("numbers", {"number": "13800000001", "department": "甲部门"}),
        ]:
            self.assertEqual(len(self.client.get(f"/api/assets/{kind}", headers=headers).json()), 1)
            self.assertEqual(
                self.client.put(f"/api/assets/{kind}/1", headers=headers, json=payload).status_code,
                403,
            )
        wb = openpyxl.Workbook()
        wb.active.append(["设备名称", "部门"])
        wb.active.append(["允许导入", "甲部门"])
        wb.active.append(["禁止导入", "乙部门"])
        data = BytesIO()
        wb.save(data)
        wb.close()
        result = self.client.post(
            "/api/ie/it/import", headers=headers, files={"file": ("scope.xlsx", data.getvalue())}
        )
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual((result.json()["imported"], result.json()["skipped"]), (1, 1))

    def test_global_read_department_write_configuration_requires_department(self):
        payload = {
            "username": "global-reader",
            "name": "跨部门查看",
            "password": "testing-password-123",
            "role": "普通用户",
            "data_scope": "all_read_dept_write",
            "permissions": ["assets", "assets_write"],
            "department": "甲部门",
        }
        result = self.client.post("/api/auth/users", headers=self.headers(1), json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["data_scope"], "all_read_dept_write")
        payload.update(username="invalid-reader", department=None)
        self.assertEqual(
            self.client.post("/api/auth/users", headers=self.headers(1), json=payload).status_code,
            400,
        )
        payload.update(department="甲部门", permissions=["departments"])
        self.assertEqual(
            self.client.post("/api/auth/users", headers=self.headers(1), json=payload).status_code,
            400,
        )

    def test_disabled_during_password_check_cannot_receive_token(self):
        entered, release = Event(), Event()

        def slow_verify(*args):
            entered.set()
            release.wait(5)
            return True

        with (
            patch("routers.auth.pwd_context.verify", side_effect=slow_verify),
            ThreadPoolExecutor(max_workers=1) as pool,
        ):
            login = pool.submit(
                self.client.post,
                "/api/auth/login",
                json={"username": "applicant", "password": "testing-password-123"},
            )
            try:
                self.assertTrue(entered.wait(2))
                result = self.client.put("/api/auth/users/2/toggle", headers=self.headers(1))
                self.assertEqual(result.status_code, 200)
            finally:
                release.set()
            self.assertEqual(login.result().status_code, 401)

    def test_login_limit_expires_and_long_password_is_rejected(self):
        with (
            patch("login_limits.monotonic", return_value=100),
            patch("routers.auth.pwd_context.verify", return_value=False),
        ):
            for _ in range(10):
                self.assertEqual(
                    self.client.post(
                        "/api/auth/login", json={"username": "admin", "password": "incorrect"}
                    ).status_code,
                    401,
                )
            self.assertEqual(
                self.client.post(
                    "/api/auth/login", json={"username": "admin", "password": "incorrect"}
                ).status_code,
                429,
            )
        with (
            patch("login_limits.monotonic", return_value=161),
            patch("routers.auth.pwd_context.verify", return_value=False),
        ):
            self.assertEqual(
                self.client.post(
                    "/api/auth/login", json={"username": "admin", "password": "incorrect"}
                ).status_code,
                401,
            )
        with patch("routers.auth.pwd_context.verify") as verify:
            self.assertEqual(
                self.client.post(
                    "/api/auth/login", json={"username": "admin", "password": "x" * 73}
                ).status_code,
                422,
            )
            verify.assert_not_called()

    def test_password_verification_does_not_hold_business_write_lock(self):
        entered, release = Event(), Event()

        def slow_verify(*args):
            entered.set()
            release.wait(5)
            return False

        with (
            patch("routers.auth.pwd_context.verify", side_effect=slow_verify),
            ThreadPoolExecutor(max_workers=2) as pool,
        ):
            login = pool.submit(
                self.client.post,
                "/api/auth/login",
                json={"username": "admin", "password": "incorrect"},
            )
            try:
                self.assertTrue(entered.wait(2))
                writing = pool.submit(
                    self.client.post,
                    "/api/assets/it",
                    headers=self.headers(1),
                    json={"name": "登录期间保存"},
                )
                self.assertEqual(writing.result(timeout=1).status_code, 200)
            finally:
                release.set()
            self.assertEqual(login.result().status_code, 401)

    def test_login_limits_attempts_and_does_not_verify_after_limit(self):
        with patch("routers.auth.pwd_context.verify", return_value=False) as verify:
            codes = [
                self.client.post(
                    "/api/auth/login", json={"username": "admin", "password": "incorrect"}
                ).status_code
                for _ in range(11)
            ]
        self.assertEqual(codes[:10], [401] * 10)
        self.assertEqual(codes[-1], 429)
        self.assertEqual(verify.call_count, 10)

    def test_excel_exports_ignore_forbidden_controls_without_changing_records(self):
        response = self.client.post(
            "/api/assets/it", headers=self.headers(1), json={"name": "bad\u0001value"}
        )
        self.assertEqual(response.status_code, 200)
        for url in ["/api/ie/it/export", "/api/reports/export"]:
            response = self.client.get(url, headers=self.headers(1))
            self.assertEqual(response.status_code, 200)
            book = openpyxl.load_workbook(BytesIO(response.content))
            self.assertTrue(
                any(cell.value == "badvalue" for sheet in book for row in sheet for cell in row)
            )
            book.close()
        self.assertEqual(
            self.client.get("/api/assets/it", headers=self.headers(1)).json()[0]["name"],
            "bad\u0001value",
        )

    def test_failed_excel_build_cleans_sheet_temporary_files(self):
        from openpyxl.worksheet._writer import ALL_TEMP_FILES

        self.asset()
        before = set(ALL_TEMP_FILES)
        with TestClient(app, raise_server_exceptions=False) as client:
            for target, url in [
                ("routers.import_export.append_text_row", "/api/ie/it/export"),
                ("routers.reports.append_text_row", "/api/reports/export"),
            ]:
                with patch(target, side_effect=RuntimeError("simulated writer failure")):
                    self.assertEqual(client.get(url, headers=self.headers(1)).status_code, 500)
                self.assertEqual(set(ALL_TEMP_FILES) - before, set())

    def test_partial_updates_preserve_fields_and_explicit_clear_is_supported(self):
        for module, identity, user_field in [
            ("it", {"name": "原名"}, "user_name"),
            ("phone", {"brand_model": "原型号"}, "user_name"),
            ("medical", {"name": "原名"}, "keeper"),
            ("numbers", {"number": "13900000001"}, "user_name"),
        ]:
            with self.subTest(module=module):
                base = "/api/assets/" + module
                original = self.client.post(
                    base,
                    headers=self.headers(1),
                    json={
                        **identity,
                        user_field: "原使用人",
                        "department": "甲部门",
                        "notes": "保留备注",
                    },
                ).json()
                response = self.client.put(
                    base + "/" + str(original["id"]),
                    headers=self.headers(1),
                    json={**identity, "notes": "新备注"},
                )
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["department"], "甲部门")
                self.assertEqual(response.json()[user_field], "原使用人")
                self.assertEqual(response.json()["status"], "in_use")
                if "asset_number" in original:
                    self.assertEqual(response.json()["asset_number"], original["asset_number"])
                response = self.client.put(
                    base + "/" + str(original["id"]),
                    headers=self.headers(1),
                    json={**identity, user_field: None},
                )
                self.assertEqual(response.status_code, 200, response.text)
                self.assertIsNone(response.json()[user_field])
                self.assertEqual(response.json()["status"], "idle")

    def test_transfer_only_account_cannot_bypass_scrap_approval(self):
        asset_id = self.asset()
        with SessionLocal() as db:
            db.get(User, 3).permissions = json.dumps(["transfer"])
            db.commit()
        response = self.client.post(
            "/api/transfer",
            headers=self.headers(3),
            json={
                "transfer_number": "ignored",
                "type": "scrap",
                "asset_desc": "资产",
                "asset_type": "it",
                "asset_id": asset_id,
            },
        )
        self.assertEqual(response.status_code, 422)
        assets = self.client.get("/api/assets/it", headers=self.headers(1)).json()
        self.assertEqual(next(a for a in assets if a["id"] == asset_id)["status"], "idle")

    def test_large_export_spools_to_disk_chunks_and_closes(self):
        import asyncio

        import excel_output

        streams = []
        original = excel_output.SpooledTemporaryFile

        def track_stream(*args, **kwargs):
            stream = original(*args, **kwargs)
            streams.append(stream)
            return stream

        class LargeWorkbook:
            def save(self, stream):
                stream.write(b"x" * (2 * 1024 * 1024 + 17))

            def close(self):
                pass

        async def consume(response):
            sizes = [len(chunk) async for chunk in response.body_iterator]
            await response.background()
            return sizes

        with patch.object(excel_output, "SpooledTemporaryFile", side_effect=track_stream):
            response = excel_output.workbook_response(LargeWorkbook(), "test.xlsx")
            self.assertTrue(streams[0]._rolled)
            sizes = asyncio.run(consume(response))
            self.assertEqual(sum(sizes), 2 * 1024 * 1024 + 17)
            self.assertLessEqual(max(sizes), 64 * 1024)
            self.assertTrue(streams[0].closed)
            response = excel_output.workbook_response(LargeWorkbook(), "test.xlsx")
            asyncio.run(response.background())
            self.assertTrue(streams[1].closed)

    def test_slow_smtp_does_not_block_other_asset_writes(self):
        self.use_medical_departments()
        entered, release = Event(), Event()
        with SessionLocal() as db:
            db.get(User, 2).email = "recipient@qq.com"
            db.add(MedicalAsset(asset_number="SMTP-1", name="测试设备", expiry_date=date.today()))
            db.commit()

        def slow_send(*args, **kwargs):
            entered.set()
            if not release.wait(8):
                raise RuntimeError("SMTP test timed out")
            return {"success": True, "message": "测试发送"}

        with (
            patch("notifier.send_email", side_effect=slow_send),
            ThreadPoolExecutor(max_workers=2) as pool,
        ):
            sending = pool.submit(
                self.client.post,
                "/api/assets/medical/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": [2]},
            )
            try:
                self.assertTrue(entered.wait(3), "SMTP was not reached")
                writing = pool.submit(
                    self.client.post,
                    "/api/assets/it",
                    headers=self.headers(1),
                    json={"name": "邮件期间正常保存"},
                )
                response = writing.result(timeout=3)
                self.assertEqual(response.status_code, 200, response.text)
            finally:
                release.set()
            self.assertEqual(sending.result(timeout=3).status_code, 200)

    def test_streamed_exports_preserve_text_dates_and_hide_passwords(self):
        with SessionLocal() as db:
            db.add(
                MedicalAsset(
                    asset_number="TEXT-1",
                    name="=1+1",
                    notes="#N/A",
                    department="=部门",
                    calibration_date=date(2026, 1, 1),
                    calibration_expiry=date(2026, 12, 31),
                )
            )
            db.add(
                WeChatAccount(
                    wx_account="export-account", wx_password=encrypt_credential("never-export-this")
                )
            )
            db.commit()
        for url in ["/api/ie/medical/export", "/api/reports/export"]:
            response = self.client.get(url, headers=self.headers(1))
            self.assertEqual(response.status_code, 200)
            with BytesIO(response.content) as source:
                book = openpyxl.load_workbook(source)
                cells = [cell for sheet in book for row in sheet for cell in row]
                literal = next(cell for cell in cells if cell.value == "=1+1")
                self.assertEqual(literal.data_type, "s")
                self.assertEqual(
                    next(cell for cell in cells if cell.value == "#N/A").data_type, "s"
                )
                self.assertTrue(any(str(cell.value).startswith("2026-12-31") for cell in cells))
                self.assertFalse(any(cell.value == "never-export-this" for cell in cells))
                self.assertFalse(any(cell.value == "微信密码" for cell in cells))
                book.close()

    def test_medical_calibration_create_update_and_optional_dates(self):
        payload = {
            "name": "校准设备",
            "department": "甲部门",
            "keeper": "保管人",
            "calibration_date": "2026-01-01",
            "calibration_expiry": "2026-12-31",
        }
        response = self.client.post("/api/assets/medical", headers=self.headers(1), json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        item_id = response.json()["id"]
        self.assertEqual(response.json()["calibration_date"], "2026-01-01")
        response = self.client.put(
            f"/api/assets/medical/{item_id}", headers=self.headers(1), json={"name": "改名"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["calibration_expiry"], "2026-12-31")
        self.assertEqual(response.json()["status"], "in_use")
        response = self.client.put(
            f"/api/assets/medical/{item_id}",
            headers=self.headers(1),
            json={"name": "改名", "calibration_expiry": "2025-01-01"},
        )
        self.assertEqual(response.status_code, 400, response.text)
        payload["calibration_expiry"] = "2025-01-01"
        self.assertEqual(
            self.client.post(
                "/api/assets/medical", headers=self.headers(1), json=payload
            ).status_code,
            422,
        )
        self.assertEqual(
            self.client.post(
                "/api/assets/medical",
                headers=self.headers(1),
                json={
                    "name": "日期可选",
                    "calibration_date": None,
                    "calibration_expiry": None,
                    "expiry_date": None,
                },
            ).status_code,
            200,
        )

    def test_calibration_reminder_boundaries_scope_and_selected_mail(self):
        self.use_medical_departments()
        today = date.today()
        with SessionLocal() as db:
            for i, (offset, status, department_id) in enumerate(
                [
                    (-1, "in_use", 1),
                    (0, "idle", 1),
                    (30, "in_use", 1),
                    (31, "in_use", 1),
                    (1, "scrapped", 1),
                    (1, "archived", 1),
                    (1, "in_use", 2),
                ]
            ):
                db.add(
                    MedicalAsset(
                        asset_number="CAL-" + str(i),
                        name="校准设备" + str(i),
                        status=status,
                        department_id=department_id,
                        calibration_date=today - timedelta(days=60),
                        calibration_expiry=today + timedelta(days=offset),
                        expiry_date=today + timedelta(days=365),
                    )
                )
            db.get(User, 2).email = "recipient@qq.com"
            db.commit()
        response = self.client.get(
            "/api/assets/medical/expiring?reminder_type=calibration", headers=self.headers(3)
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual([r["days_left"] for r in response.json()], [-1, 0, 30])
        self.assertEqual(
            self.client.get("/api/assets/medical/expiring", headers=self.headers(3)).json(), []
        )
        with patch(
            "notifier.send_email", return_value={"success": True, "message": "测试发送"}
        ) as mock:
            response = self.client.post(
                "/api/assets/medical/send-expiry-email?reminder_type=calibration",
                headers=self.headers(1),
                json={"recipient_user_ids": [2]},
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["count"], 4)
            self.assertEqual(mock.call_args.args[0], ["recipient@qq.com"])
            self.assertIn("校准到期提醒", mock.call_args.args[1])
            self.assertIn("校准时间", mock.call_args.args[2])
            self.assertNotIn("CAL-3", mock.call_args.args[2])
        self.assertEqual(
            self.client.post(
                "/api/assets/medical/send-expiry-email?reminder_type=calibration",
                headers=self.headers(3),
                json={"recipient_user_ids": [2]},
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.get(
                "/api/assets/medical/expiring?reminder_type=other", headers=self.headers(1)
            ).status_code,
            422,
        )
        self.assertEqual(
            self.client.get(
                "/api/assets/medical/expiring?days=0", headers=self.headers(1)
            ).status_code,
            400,
        )

    def test_calibration_empty_does_not_send_mail(self):
        from notifier import send_expiry_notice

        with SessionLocal() as db, patch("notifier.send_email") as mock:
            response = send_expiry_notice(db, 30, ["recipient@qq.com"], "calibration")
            self.assertEqual(response["count"], 0)
            mock.assert_not_called()

    def test_medical_calibration_excel_new_and_legacy_headers(self):
        def workbook(headers, rows):
            book = openpyxl.Workbook()
            book.active.append(headers)
            for row in rows:
                book.active.append(row)
            buf = BytesIO()
            book.save(buf)
            book.close()
            return buf.getvalue()

        headers = ["设备名称", "校准时间", "校准到期时间"]
        content = workbook(
            headers,
            [["有效校准", "2026-01-01", "2026-12-31"], ["错误日期", "2026-12-31", "2026-01-01"]],
        )
        response = self.client.post(
            "/api/ie/medical/import",
            headers=self.headers(1),
            files={"file": ("medical.xlsx", content)},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["imported"], 1)
        self.assertIn("不能早于", response.json()["errors"][0])
        legacy = workbook(["设备名称", "校准到期"], [["旧模板", "2027-01-01"]])
        self.assertEqual(
            self.client.post(
                "/api/ie/medical/import",
                headers=self.headers(1),
                files={"file": ("medical.xlsx", legacy)},
            ).json()["imported"],
            1,
        )
        response = self.client.get("/api/ie/medical/export", headers=self.headers(1))
        book = openpyxl.load_workbook(BytesIO(response.content))
        rows = list(book.active.values)
        book.close()
        self.assertIn("校准时间", rows[0])
        self.assertIn("校准到期时间", rows[0])
        self.assertEqual(str(rows[1][rows[0].index("校准时间")])[:10], "2026-01-01")

    def user_workbook(self, data_rows, active_notes=False):
        from routers.user_import import HEADERS

        wb = openpyxl.Workbook()
        wb.active.append(HEADERS)
        for row in data_rows:
            wb.active.append(row)
        if active_notes:
            wb.create_sheet("说明").append(["说明"])
            wb.active = 1
        buffer = BytesIO()
        wb.save(buffer)
        wb.close()
        return buffer.getvalue()

    def test_user_import_template_superadmin_only(self):
        response = self.client.get("/api/user-import/template", headers=self.headers(1))
        self.assertEqual(response.status_code, 200, response.text)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["新增用户", "填写说明", "权限对照"])
        self.assertEqual(wb.worksheets[0]["A1"].value, "用户名")
        self.assertIsNone(wb.worksheets[0]["C2"].value)
        wb.close()
        self.assertEqual(
            self.client.get("/api/user-import/template", headers=self.headers(3)).status_code, 403
        )
        self.assertEqual(
            self.client.post(
                "/api/user-import/import",
                headers=self.headers(3),
                files={"file": ("users.xlsx", self.user_workbook([]))},
            ).status_code,
            403,
        )

    def test_user_import_creates_independent_permissions_and_valid_credentials(self):
        content = self.user_workbook(
            [
                [
                    "newreader",
                    "新用户",
                    "reader-password-123",
                    "new@qq.com",
                    "自由角色",
                    "甲部门",
                    "仅本部门",
                    "查看资产档案、导出数据",
                ],
                [
                    "newblank",
                    "无权限用户",
                    "blank-password-123",
                    "",
                    "另一角色",
                    "",
                    "全部部门",
                    "",
                ],
            ],
            active_notes=True,
        )
        response = self.client.post(
            "/api/user-import/import",
            headers=self.headers(1),
            files={"file": ("users.xlsx", content)},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["count"], 2)
        with SessionLocal() as db:
            user = db.query(User).filter(User.username == "newreader").one()
            self.assertTrue(pwd_context.verify("reader-password-123", user.password_hash))
            self.assertNotEqual(user.password_hash, "reader-password-123")
            self.assertEqual(user.department_id, 1)
            self.assertEqual(user.email, "new@qq.com")
            self.assertIsNone(user.role_id)
            self.assertEqual(json.loads(user.permissions), ["assets", "export"])
            self.assertEqual(
                db.query(User).filter(User.username == "newblank").one().permissions, "[]"
            )
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "newreader", "password": "reader-password-123"}
            ).status_code,
            200,
        )

    def test_user_import_errors_are_atomic_and_never_echo_passwords(self):
        content = self.user_workbook(
            [
                [
                    "wouldpass",
                    "有效用户",
                    "valid-password-123",
                    "",
                    "岗位",
                    "",
                    "全部部门",
                    "assets",
                ],
                [
                    "broken",
                    "无效用户",
                    "shortSECRET",
                    "bademail",
                    "super_admin",
                    "",
                    "全部部门",
                    "users",
                ],
            ]
        )
        response = self.client.post(
            "/api/user-import/import",
            headers=self.headers(1),
            files={"file": ("users.xlsx", content)},
        )
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(response.json()["detail"]["errors"][0]["row"], 3)
        self.assertNotIn("shortSECRET", response.text)
        with SessionLocal() as db:
            self.assertEqual(db.query(User).count(), 5)
        for rows in [
            [["admin", "重复", "testing-password-123", "", "", "", "全部部门", ""]],
            [
                ["dupe", "重复1", "testing-password-123", "", "", "", "全部部门", ""],
                ["dupe", "重复2", "testing-password-123", "", "", "", "全部部门", ""],
            ],
            [["badmail", "坏邮箱", "testing-password-123", "bademail", "", "", "全部部门", ""]],
            [["badgrant", "越权", "testing-password-123", "", "", "甲部门", "仅本部门", "logs"]],
        ]:
            response = self.client.post(
                "/api/user-import/import",
                headers=self.headers(1),
                files={"file": ("users.xlsx", self.user_workbook(rows))},
            )
            self.assertEqual(response.status_code, 400, response.text)
            with SessionLocal() as db:
                self.assertEqual(db.query(User).count(), 5)

    def test_user_import_rejects_empty_invalid_and_excessive_files(self):
        for filename, content in [
            ("users.xlsx", self.user_workbook([])),
            ("users.csv", b"csv"),
            ("users.xlsx", b"not xlsx"),
            (
                "users.xlsx",
                self.user_workbook(
                    [
                        ["x" + str(i), "测试", "testing-password-123", "", "", "", "全部部门", ""]
                        for i in range(101)
                    ]
                ),
            ),
        ]:
            response = self.client.post(
                "/api/user-import/import",
                headers=self.headers(1),
                files={"file": (filename, content)},
            )
            self.assertEqual(response.status_code, 400, response.text)

    def test_existing_admin_empty_departments_remain_empty_on_restart(self):
        from main import seed_admin

        with SessionLocal() as db:
            db.execute(text("DELETE FROM users WHERE id <> 1"))
            db.execute(
                text("UPDATE users SET department=NULL,department_id=NULL,role_id=NULL WHERE id=1")
            )
            db.execute(text("DELETE FROM departments"))
            db.commit()
        seed_admin()
        with SessionLocal() as db:
            self.assertEqual(db.query(User).count(), 1)
            self.assertEqual(db.query(Department).count(), 0)

    @classmethod
    def setUpClass(cls):
        cls.password_hash = pwd_context.hash("testing-password-123")

    def setUp(self):
        from login_limits import login_limiter, password_limiter

        login_limiter.clear()
        password_limiter.clear()
        with SessionLocal() as db:
            for table in reversed(Base.metadata.sorted_tables):
                db.execute(table.delete())
            db.add_all([Department(id=1, name="甲部门"), Department(id=2, name="乙部门")])
            db.add_all(
                [
                    User(
                        id=1,
                        username="admin",
                        name="管理员",
                        role="super_admin",
                        password_hash=self.password_hash,
                        data_scope="all",
                    ),
                    User(
                        id=2,
                        username="applicant",
                        name="同名",
                        role="",
                        password_hash=self.password_hash,
                        department="甲部门",
                        department_id=1,
                        permissions=json.dumps(["apply"]),
                        data_scope="department",
                    ),
                    User(
                        id=3,
                        username="approver",
                        name="同名",
                        role="",
                        password_hash=self.password_hash,
                        department="甲部门",
                        department_id=1,
                        permissions=json.dumps(["approval", "transfer", "assets_write"]),
                        data_scope="department",
                    ),
                    User(
                        id=4,
                        username="foreign",
                        name="乙用户",
                        role="",
                        password_hash=self.password_hash,
                        department="乙部门",
                        department_id=2,
                        permissions=json.dumps(["apply", "approval"]),
                        data_scope="department",
                    ),
                    User(
                        id=5,
                        username="report",
                        name="报表员",
                        role="",
                        password_hash=self.password_hash,
                        department="甲部门",
                        department_id=1,
                        permissions=json.dumps(["reports"]),
                        data_scope="department",
                    ),
                ]
            )
            db.commit()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()

    def headers(self, user_id):
        with SessionLocal() as db:
            u = db.get(User, user_id)
            return {
                "Authorization": "Bearer " + create_token(u.username, u.token_version or 0, u.id)
            }

    def test_user_email_create_update_clear_and_legacy_preserve(self):
        payload = {
            "username": "mailuser",
            "password": "testing-password-123",
            "name": "邮箱用户",
            "role": "",
            "department": "甲部门",
            "data_scope": "department",
            "permissions": [],
            "email": " test@qq.com ",
        }
        response = self.client.post("/api/auth/users", headers=self.headers(1), json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["email"], "test@qq.com")
        user_id = response.json()["id"]
        payload.pop("email")
        payload.pop("password")
        response = self.client.put(
            f"/api/auth/users/{user_id}", headers=self.headers(1), json=payload
        )
        self.assertEqual(response.json()["email"], "test@qq.com")
        payload["email"] = "not-an-email"
        self.assertEqual(
            self.client.put(
                f"/api/auth/users/{user_id}", headers=self.headers(1), json=payload
            ).status_code,
            422,
        )
        payload["email"] = ""
        self.assertIsNone(
            self.client.put(
                f"/api/auth/users/{user_id}", headers=self.headers(1), json=payload
            ).json()["email"]
        )

    def test_expiry_recipient_permissions_and_selected_delivery(self):
        self.use_medical_departments()
        with SessionLocal() as db:
            db.get(User, 2).email = "selected@qq.com"
            db.get(User, 4).email = "unselected@qq.com"
            db.add(
                MedicalAsset(
                    asset_number="MAIL-TEST",
                    name="<script>测试</script>",
                    expiry_date=date.today(),
                    status="in_use",
                )
            )
            db.commit()
        self.assertEqual(
            self.client.get(
                "/api/assets/medical/notice-recipients", headers=self.headers(3)
            ).status_code,
            403,
        )
        self.assertEqual(self.client.get("/api/assets/medical/notice-recipients").status_code, 403)
        response = self.client.get("/api/assets/medical/notice-recipients", headers=self.headers(1))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual({u["id"] for u in response.json()["users"]}, {2, 4})
        with patch(
            "notifier.send_email", return_value={"success": True, "message": "已发送"}
        ) as mock:
            response = self.client.post(
                "/api/assets/medical/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": [2, 2]},
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(mock.call_args.args[0], ["selected@qq.com"])
            self.assertIn("&lt;script&gt;", mock.call_args.args[2])
            self.assertNotIn("<script>", mock.call_args.args[2])
        with SessionLocal() as db:
            db.get(User, 2).is_active = 0
            db.commit()
        with patch("notifier.send_email") as mock:
            response = self.client.post(
                "/api/assets/medical/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": [2]},
            )
            self.assertEqual(response.status_code, 400)
            mock.assert_not_called()
        self.assertEqual(
            self.client.post(
                "/api/assets/medical/send-expiry-email",
                headers=self.headers(1),
                json={"recipient_user_ids": []},
            ).status_code,
            422,
        )

    def asset(self, department_id=1, status="idle", number=None):
        with SessionLocal() as db:
            asset = ITAsset(
                asset_number=number or f"TEST-{department_id}",
                name="测试设备",
                department_id=department_id,
                department="甲部门" if department_id == 1 else "乙部门",
                status=status,
            )
            db.add(asset)
            db.commit()
            return asset.id

    def apply(self, asset_id, user_id=2, action="checkout"):
        return self.client.post(
            "/api/delete-request",
            headers=self.headers(user_id),
            json={
                "table_name": "apply_requests",
                "record_desc": "测试申请",
                "reason": "用途中包含|也应正常",
                "application_type": action,
                "asset_type": "it",
                "asset_id": asset_id,
                "target_user": "接收人",
                "target_department_id": 1 if user_id != 4 else 2,
            },
        )

    def test_role_label_never_assigns_permissions(self):
        with SessionLocal() as db:
            db.add(Role(name="只读", permissions=json.dumps(["assets", "assets_write"])))
            db.commit()
        response = self.client.post(
            "/api/auth/users",
            headers=self.headers(1),
            json={
                "username": "reader",
                "password": "reader-password-123",
                "name": "只读用户",
                "role": "只读",
                "department": "甲部门",
                "data_scope": "department",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["permissions"], [])
        self.assertIsNone(response.json()["role_id"])
        user_id = response.json()["id"]
        payload = {
            "username": "reader",
            "name": "只读用户",
            "role": "自定义岗位",
            "department": "甲部门",
            "data_scope": "department",
            "permissions": ["assets"],
        }
        response = self.client.put(
            f"/api/auth/users/{user_id}", headers=self.headers(1), json=payload
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["role"], "自定义岗位")
        self.assertEqual(response.json()["permissions"], ["assets"])
        payload.pop("permissions")
        payload["role"] = "修改名称"
        self.assertEqual(
            self.client.put(
                f"/api/auth/users/{user_id}", headers=self.headers(1), json=payload
            ).json()["permissions"],
            ["assets"],
        )

    def test_retired_roles_and_superadmin_only_assignment(self):
        for method, path in [
            ("get", "/api/auth/roles"),
            ("post", "/api/auth/roles"),
            ("put", "/api/auth/roles/1"),
            ("delete", "/api/auth/roles/1"),
        ]:
            self.assertEqual(
                getattr(self.client, method)(path, headers=self.headers(1)).status_code, 410
            )
            self.assertEqual(
                getattr(self.client, method)(path, headers=self.headers(3)).status_code, 403
            )
        self.assertNotIn(
            "roles", self.client.get("/api/auth/me", headers=self.headers(1)).json()["permissions"]
        )
        payload = {
            "username": "applicant",
            "name": "同名",
            "role": "超级管理员",
            "department": "甲部门",
            "data_scope": "all",
            "permissions": ["assets"],
        }
        self.assertEqual(
            self.client.put("/api/auth/users/2", headers=self.headers(3), json=payload).status_code,
            403,
        )
        self.assertEqual(
            self.client.put("/api/auth/users/2", headers=self.headers(1), json=payload).status_code,
            200,
        )
        self.assertEqual(
            self.client.get("/api/auth/users", headers=self.headers(2)).status_code, 403
        )

    def test_legacy_template_is_snapshotted_once_then_decoupled(self):
        from deps import get_user_permissions
        from main import seed_admin

        with SessionLocal() as db:
            legacy = Role(name="旧角色", permissions=json.dumps(["assets"]))
            db.add(legacy)
            db.flush()
            user = db.get(User, 2)
            user.role = "旧角色"
            user.role_id = legacy.id
            user.permissions = None
            db.commit()
            self.assertEqual(get_user_permissions(user, db), [])
        seed_admin()
        with SessionLocal() as db:
            user = db.get(User, 2)
            self.assertEqual(get_user_permissions(user, db), ["assets"])
            db.get(Role, user.role_id).permissions = json.dumps(["assets", "assets_write"])
            db.commit()
            self.assertEqual(get_user_permissions(user, db), ["assets"])

    def test_global_permissions_rejected_for_department_scope(self):
        for permission in ["departments", "logs", "wechat"]:
            with SessionLocal() as db:
                db.get(User, 2).permissions = json.dumps([permission])
                db.commit()
            endpoint = (
                "/api/logs"
                if permission == "logs"
                else "/api/wechat"
                if permission == "wechat"
                else "/api/departments"
            )
            response = (
                self.client.get(endpoint, headers=self.headers(2))
                if permission != "departments"
                else self.client.put(
                    endpoint + "/2", headers=self.headers(2), json={"name": "篡改"}
                )
            )
            self.assertEqual(response.status_code, 403, response.text)
        response = self.client.put(
            "/api/auth/users/2",
            headers=self.headers(1),
            json={
                "username": "applicant",
                "name": "同名",
                "role": "",
                "department": "甲部门",
                "permissions": ["logs"],
                "data_scope": "department",
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_cross_department_approve_and_reject_denied(self):
        asset = self.asset(2)
        response = self.apply(asset, user_id=4)
        self.assertEqual(response.status_code, 200, response.text)
        for action in ["approve", "reject"]:
            result = self.client.put(
                f"/api/delete-request/{response.json()['id']}/{action}", headers=self.headers(3)
            )
            self.assertEqual(result.status_code, 403, result.text)

    def test_scrapped_asset_cannot_be_resurrected(self):
        asset = self.asset()
        req = self.apply(asset)
        self.assertEqual(req.status_code, 200, req.text)
        with SessionLocal() as db:
            db.get(ITAsset, asset).status = "scrapped"
            db.commit()
        result = self.client.put(
            f"/api/delete-request/{req.json()['id']}/approve", headers=self.headers(3)
        )
        self.assertEqual(result.status_code, 409, result.text)
        with SessionLocal() as db:
            self.assertEqual(db.get(ITAsset, asset).status, "scrapped")
            self.assertEqual(db.get(DeleteRequest, req.json()["id"]).status, "pending")

    def test_same_name_accounts_have_distinct_ownership(self):
        asset = self.asset()
        req = self.apply(asset)
        self.assertEqual(req.status_code, 200, req.text)
        result = self.client.put(
            f"/api/delete-request/{req.json()['id']}/approve", headers=self.headers(3)
        )
        self.assertEqual(result.status_code, 200, result.text)
        with SessionLocal() as db:
            self.assertEqual(db.get(DeleteRequest, req.json()["id"]).approver_id, 3)

    def test_self_approval_denied_by_id_after_name_change(self):
        asset = self.asset()
        req = self.apply(asset)
        with SessionLocal() as db:
            u = db.get(User, 2)
            u.name = "新姓名"
            u.permissions = json.dumps(["approval"])
            db.commit()
        result = self.client.put(
            f"/api/delete-request/{req.json()['id']}/approve", headers=self.headers(2)
        )
        self.assertEqual(result.status_code, 400, result.text)

    def test_concurrent_approval_applies_once(self):
        req = self.apply(self.asset()).json()
        headers = self.headers(3)

        def approve(_):
            with TestClient(app) as client:
                return client.put(
                    f"/api/delete-request/{req['id']}/approve", headers=headers
                ).status_code

        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(approve, range(2))), [200, 409])
        with SessionLocal() as db:
            self.assertEqual(db.query(TransferRecord).count(), 1)

    def test_numbering_survives_deletion_and_manual_high_number(self):
        stem = f"IT-{date.today().year}-"
        self.asset(number=stem + "001")
        self.asset(number=stem + "003")
        payload = {"name": "新增", "department": "甲部门"}
        result = self.client.post("/api/assets/it", headers=self.headers(3), json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["asset_number"], stem + "004")
        with SessionLocal() as db:
            db.delete(db.get(ITAsset, result.json()["id"]))
            db.commit()
        result = self.client.post("/api/assets/it", headers=self.headers(3), json=payload)
        self.assertEqual(result.json()["asset_number"], stem + "005")
        self.asset(number=stem + "099")
        result = self.client.post("/api/assets/it", headers=self.headers(3), json=payload)
        self.assertEqual(result.json()["asset_number"], stem + "100")

    def test_concurrent_number_generation_is_unique(self):
        headers = self.headers(3)

        def create(i):
            response = self.client.post(
                "/api/assets/it", headers=headers, json={"name": f"设备{i}", "department": "甲部门"}
            )
            self.assertEqual(response.status_code, 200, response.text)
            return response.json()["asset_number"]

        with ThreadPoolExecutor(max_workers=4) as pool:
            numbers = list(pool.map(create, range(6)))
        self.assertEqual(len(set(numbers)), 6)

    def test_department_rename_preserves_pending_and_history(self):
        asset = self.asset()
        req = self.client.post(
            "/api/scrap",
            headers=self.headers(2),
            json={"asset_desc": "测试", "asset_type": "it", "asset_id": asset},
        )
        self.assertEqual(req.status_code, 200, req.text)
        rename = self.client.put(
            "/api/departments/1", headers=self.headers(1), json={"name": "甲部门新名"}
        )
        self.assertEqual(rename.status_code, 200, rename.text)
        result = self.client.put(f"/api/scrap/{req.json()['id']}/approve", headers=self.headers(3))
        self.assertEqual(result.status_code, 200, result.text)
        history = self.client.get("/api/transfer", headers=self.headers(3)).json()
        self.assertEqual(history[0]["department_id"], 1)
        self.assertEqual(history[0]["department"], "甲部门新名")
        self.assertEqual(history[0]["asset_id"], asset)
        delete = self.client.delete("/api/departments/1", headers=self.headers(1))
        self.assertEqual(delete.status_code, 409)

    def test_transfer_records_before_and_after(self):
        asset = self.asset(status="in_use")
        with SessionLocal() as db:
            db.get(ITAsset, asset).user_name = "原使用人"
            db.commit()
        result = self.client.post(
            "/api/transfer",
            headers=self.headers(1),
            json={
                "transfer_number": "ignored-client-number",
                "type": "transfer",
                "asset_desc": "测试",
                "asset_type": "it",
                "asset_id": asset,
                "new_user": "新使用人",
                "new_dept": "乙部门",
            },
        )
        self.assertEqual(result.status_code, 200, result.text)
        data = result.json()
        self.assertEqual(
            (data["department_id"], data["new_department_id"], data["old_user"], data["new_user"]),
            (1, 2, "原使用人", "新使用人"),
        )
        with SessionLocal() as db:
            db.get(User, 4).permissions = json.dumps(["transfer"])
            db.commit()
        self.assertEqual(len(self.client.get("/api/transfer", headers=self.headers(4)).json()), 1)

    def test_asset_removal_preserves_identity_and_hides_archive(self):
        asset = self.asset()
        request = self.client.post(
            "/api/delete-request",
            headers=self.headers(3),
            json={
                "table_name": "it_assets",
                "record_id": asset,
                "record_desc": "移出档案",
                "reason": "重复档案",
            },
        )
        self.assertEqual(request.status_code, 200, request.text)
        approved = self.client.put(
            f"/api/delete-request/{request.json()['id']}/approve", headers=self.headers(1)
        )
        self.assertEqual(approved.status_code, 200, approved.text)
        with SessionLocal() as db:
            self.assertEqual(db.get(ITAsset, asset).status, "archived")
        self.assertEqual(self.client.get("/api/assets/it", headers=self.headers(3)).json(), [])
        edit = self.client.put(
            f"/api/assets/it/{asset}",
            headers=self.headers(3),
            json={"name": "恢复", "department": "甲部门"},
        )
        self.assertEqual(edit.status_code, 404)
        request = self.client.post(
            "/api/delete-request",
            headers=self.headers(3),
            json={
                "table_name": "it_assets",
                "record_id": asset,
                "record_desc": "再次删除",
                "reason": "重复",
            },
        )
        self.assertEqual(request.status_code, 409)

    def test_user_with_history_cannot_be_deleted(self):
        self.apply(self.asset())
        result = self.client.delete("/api/auth/users/2", headers=self.headers(1))
        self.assertEqual(result.status_code, 409, result.text)

    def test_superadmin_can_save_existing_full_permission_list(self):
        me = self.client.get("/api/auth/me", headers=self.headers(1)).json()
        me["password"] = ""
        result = self.client.put("/api/auth/users/1", headers=self.headers(1), json=me)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["role"], "super_admin")

    def test_password_reset_and_toggle_revoke_old_tokens(self):
        old = self.headers(2)
        update = self.client.put(
            "/api/auth/users/2",
            headers=self.headers(1),
            json={
                "username": "applicant",
                "name": "同名",
                "role": "",
                "department": "甲部门",
                "permissions": ["apply"],
                "data_scope": "department",
                "password": "changed-password-123",
            },
        )
        self.assertEqual(update.status_code, 200, update.text)
        self.assertEqual(self.client.get("/api/auth/me", headers=old).status_code, 401)
        fresh = self.headers(2)
        self.client.put("/api/auth/users/2/toggle", headers=self.headers(1))
        self.client.put("/api/auth/users/2/toggle", headers=self.headers(1))
        self.assertEqual(self.client.get("/api/auth/me", headers=fresh).status_code, 401)

    def test_login_and_recreated_account_do_not_restore_deleted_session(self):
        login = self.client.post(
            "/api/auth/login", json={"username": "report", "password": "testing-password-123"}
        )
        self.assertEqual(login.status_code, 200, login.text)
        old = {"Authorization": "Bearer " + login.json()["token"]}
        self.assertEqual(self.client.get("/api/auth/me", headers=old).status_code, 200)
        deleted = self.client.delete("/api/auth/users/5", headers=self.headers(1))
        self.assertEqual(deleted.status_code, 200, deleted.text)
        recreated = self.client.post(
            "/api/auth/users",
            headers=self.headers(1),
            json={
                "username": "report",
                "name": "新账号",
                "role": "",
                "department": "甲部门",
                "data_scope": "department",
                "permissions": ["reports"],
                "password": "new-password-123",
            },
        )
        self.assertEqual(recreated.status_code, 200, recreated.text)
        self.assertEqual(self.client.get("/api/auth/me", headers=old).status_code, 401)

    def test_reports_export_permission_scope_and_no_wechat(self):
        self.asset(1)
        self.asset(2)
        self.assertEqual(
            self.client.get("/api/reports/export", headers=self.headers(5)).status_code, 403
        )
        with SessionLocal() as db:
            db.get(User, 5).permissions = json.dumps(["reports", "export"])
            db.commit()
        result = self.client.get("/api/reports/export", headers=self.headers(5))
        self.assertEqual(result.status_code, 200, result.text if result.status_code != 200 else "")
        wb = openpyxl.load_workbook(BytesIO(result.content))
        self.assertNotIn("微信账号", wb.sheetnames)
        rows = list(wb[wb.sheetnames[0]].values)
        self.assertEqual(len(rows), 2)
        self.assertIn("TEST-1", rows[1])
        self.assertNotIn("TEST-2", rows[1])
        wb.close()

    def test_password_prefix_and_corrupt_row_are_isolated(self):
        value = "enc:v1:hello"
        self.assertEqual(decrypt_credential(encrypt_credential(value)), value)
        with SessionLocal() as db:
            db.add_all(
                [
                    WeChatAccount(wx_account="valid", wx_password=encrypt_credential(value)),
                    WeChatAccount(wx_account="broken", wx_password="enc:v1:bad"),
                ]
            )
            db.commit()
        result = self.client.get("/api/wechat", headers=self.headers(1))
        self.assertEqual(result.status_code, 200, result.text)
        bad = next(row for row in result.json() if row["wx_account"] == "broken")
        self.assertIsNone(bad["wx_password"])
        self.assertIn("password_error", bad)

    def test_import_uses_shared_sequence_and_rejects_invalid_dates(self):
        self.asset(number=f"IT-{date.today().year}-003")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["设备名称", "部门", "购入日期"])
        ws.append(["正确", "甲部门", "2026-10-09"])
        ws.append(["错误", "甲部门", "不是日期"])
        data = BytesIO()
        wb.save(data)
        wb.close()
        result = self.client.post(
            "/api/ie/it/import",
            headers=self.headers(1),
            files={"file": ("test.xlsx", data.getvalue())},
        )
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual((result.json()["imported"], result.json()["skipped"]), (1, 1))
        self.assertIn("日期", result.json()["errors"][0])
        with SessionLocal() as db:
            self.assertEqual(
                db.query(ITAsset).filter(ITAsset.name == "正确").one().asset_number,
                f"IT-{date.today().year}-004",
            )

    def test_import_rejects_excess_rows_and_upload_size(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(10002, 1, "过多行")
        data = BytesIO()
        wb.save(data)
        wb.close()
        result = self.client.post(
            "/api/ie/it/import",
            headers=self.headers(1),
            files={"file": ("test.xlsx", data.getvalue())},
        )
        self.assertEqual(result.status_code, 413, result.text)
        result = self.client.post(
            "/api/ie/it/import",
            headers=self.headers(1),
            files={"file": ("test.xlsx", b"x" * (10 * 1024 * 1024 + 1))},
        )
        self.assertEqual(result.status_code, 413, result.text)


class MigrationTests(unittest.TestCase):
    def test_legacy_ids_are_unique_and_migration_runs_once(self):
        other = create_engine(
            "sqlite:///" + str(Path(workspace.name) / "migration.db").replace("\\", "/")
        )
        with other.begin() as conn:
            conn.execute(text("CREATE TABLE departments(id INTEGER PRIMARY KEY, name TEXT)"))
            conn.execute(
                text(
                    "CREATE TABLE users(id INTEGER PRIMARY KEY,username TEXT,name TEXT,role TEXT,password_hash TEXT,data_scope TEXT)"
                )
            )
            conn.execute(
                text(
                    "CREATE TABLE transfer_records(id INTEGER PRIMARY KEY, department TEXT,new_dept TEXT,operator TEXT)"
                )
            )
            conn.execute(
                text(
                    "CREATE TABLE scrap_requests(id INTEGER PRIMARY KEY,request_number TEXT,asset_desc TEXT,applicant TEXT,department TEXT,auditor TEXT)"
                )
            )
            conn.execute(
                text(
                    "CREATE TABLE delete_requests(id INTEGER PRIMARY KEY,table_name TEXT,record_id INTEGER,reason TEXT,applicant TEXT,approver TEXT)"
                )
            )
            for table in ["it_assets", "phone_assets", "medical_assets", "phone_numbers"]:
                conn.execute(
                    text(
                        f"CREATE TABLE {table}(id INTEGER PRIMARY KEY,asset_number TEXT,name TEXT,department TEXT,department_id INTEGER,status TEXT)"
                    )
                )
            conn.execute(text("INSERT INTO departments(id,name) VALUES(1,'旧部门')"))
            conn.execute(
                text(
                    "INSERT INTO users(id,username,name,role,password_hash,data_scope) VALUES(1,'one','重名','','x','all'),(2,'two','重名','','x','all'),(3,'unique','唯一','','x','all')"
                )
            )
            conn.execute(
                text(
                    "INSERT INTO it_assets(id,asset_number,name,department,department_id,status) VALUES(1,'IT-1','测试','旧部门',1,'idle')"
                )
            )
            conn.execute(
                text(
                    "INSERT INTO scrap_requests(request_number,asset_desc,applicant,department) VALUES('SC-1','测试','重名','旧部门'),('SC-2','测试','唯一','旧部门')"
                )
            )
            conn.execute(
                text(
                    "INSERT INTO delete_requests(table_name,record_id,reason,applicant) VALUES('apply_requests',0,'checkout|接收人|旧部门|it|1','唯一')"
                )
            )
        migrate_workflow(other)
        with other.begin() as conn:
            rows = conn.execute(
                text("SELECT applicant_id,department_id FROM scrap_requests ORDER BY id")
            ).all()
            self.assertEqual(rows, [(None, 1), (3, 1)])
            request = conn.execute(
                text(
                    "SELECT application_type,target_department_id,department_id,asset_id FROM delete_requests"
                )
            ).one()
            self.assertEqual(tuple(request), ("checkout", 1, 1, 1))
            conn.execute(text("DELETE FROM users WHERE id=2"))
        migrate_workflow(other)
        with other.connect() as conn:
            self.assertIsNone(
                conn.execute(text("SELECT applicant_id FROM scrap_requests WHERE id=1")).scalar()
            )
        other.dispose()


def tearDownModule():
    engine.dispose()
    workspace.cleanup()


if __name__ == "__main__":
    unittest.main()
