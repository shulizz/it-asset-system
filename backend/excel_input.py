"""Bounded XLSX validation and parsing, performed outside the event loop."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import iterparse
from fastapi import HTTPException

MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
MAX_ROWS = 10001  # header + 10,000 data rows
MAX_COLUMNS = 100


def read_rows(content):
    import openpyxl
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(f.file_size for f in archive.infolist()) > MAX_EXPANDED:
                raise HTTPException(413, 'Excel解压后过大或文件条目过多')
            expanded = 0
            for info in archive.infolist():
                with archive.open(info) as entry:
                    while chunk := entry.read(1024 * 1024):
                        expanded += len(chunk)
                        if expanded > MAX_EXPANDED:
                            raise HTTPException(413, 'Excel解压后不能超过50MB')
                if info.filename.startswith('xl/worksheets/') and info.filename.endswith('.xml'):
                    from openpyxl.utils.cell import coordinate_to_tuple
                    with archive.open(info) as sheet:
                        for _, element in iterparse(sheet, events=('start',)):
                            tag = element.tag.rsplit('}', 1)[-1]
                            if tag == 'row' and int(element.attrib.get('r', '0')) > MAX_ROWS:
                                raise HTTPException(413, '最多导入10000条数据')
                            if tag == 'c' and element.attrib.get('r'):
                                row, column = coordinate_to_tuple(element.attrib['r'])
                                if row > MAX_ROWS or column > MAX_COLUMNS:
                                    raise HTTPException(413, '最多导入10000条数据，最多100列')
                            element.clear()
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise HTTPException(400, '文件没有可读取的工作表')
            if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLUMNS:
                raise HTTPException(413, '最多导入10000条数据，最多100列')
            rows = []
            # Ignore forged dimensions; iterate actual cells with hard limits.
            ws.reset_dimensions()
            for row in ws.iter_rows(values_only=True, max_col=MAX_COLUMNS):
                if len(rows) >= MAX_ROWS:
                    raise HTTPException(413, '最多导入10000条数据，最多100列')
                rows.append(row)
            return rows
        finally:
            wb.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, '无法读取文件，请上传有效的.xlsx文件')
